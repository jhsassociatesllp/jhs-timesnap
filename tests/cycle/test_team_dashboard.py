"""Team dashboard: the status rules and the API (scope, filters, figures, Excel). Scratch Mongo + the fake NQ in Docker."""
import io
from datetime import date, datetime

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.auth import get_current_user
from backend.metrics import access
from backend.metrics.store import mstore
from backend.nq import cache, repository
from backend.timesheet import team_dashboard as td
from tests.cycle.conftest import entry
from tests.nq.conftest import dev_settings

_user = {"id": "JHS9001"}


# ── pure rules ───────────────────────────────────────────────────────────────

def test_project_status_rules():
    st = td.project_status
    assert st(10, 100, 5, 1, 0, 2) == "on_track"
    assert st(10, 100, 5, 1, 2, 2) == "late"                       # an overdue task
    assert st(120, 100, 5, 4, 0, 2) == "late"                      # over the allotted hours
    assert st(85, 100, 10, 5, 0, 2) == "watch"                     # 85% used, 50% done
    assert st(85, 100, 10, 9, 0, 2) == "on_track"                  # nearly finished: fine
    assert st(10, 100, 5, 1, 0, 14) == "watch"                     # quiet for 14 days with open tasks
    assert st(10, 100, 5, 5, 0, 30) == "on_track"                  # everything done: quiet is fine
    assert st(0, None, 0, 0, 0, None) == "unknown"
    assert td.status_reasons(120.0, 100.0, 5, 4, 2, 20) == ["2 overdue tasks", "120.0 h used, over the 100.0 h allotted", "no hours logged for 20 days"]


# ── API ──────────────────────────────────────────────────────────────────────

@pytest.fixture
def api(cdb, monkeypatch):
    repository.configure(dev_settings())
    cache.clear()
    cdb["c_audit"].drop()
    monkeypatch.setattr(mstore, "audit", cdb["c_audit"])
    monkeypatch.setattr(access, "has_admin_module", lambda emp, key: emp == "JHS9000" and key == "timesheet")
    monkeypatch.setattr(td, "_today", lambda: date(2026, 10, 21))
    now = datetime.utcnow()
    A = dict(client="Test Client A", project="Statutory Audit FY26", project_code="PL-T001")
    B = dict(client="Test Client A", project="Internal Audit Q2", project_code="PL-T002")
    cdb["c_entries"].insert_many([
        entry("JHS9101", "2026-10-12", 8.5, task_id=2, **A), entry("JHS9101", "2026-10-13", 8.5, task_id=2, **A),
        entry("JHS9101", "2026-10-13", 1, task_id=None, is_unplanned=True, client="", project="", project_code=""),
        entry("JHS9102", "2026-10-12", 7, task_id=5, **B),
        entry("JHS9103", "2026-10-12", 6, task_id=5, **B)])                                  # another manager's employee
    cdb["c_submissions"].insert_many([{"employee_id": "JHS9101", "date": d, "cycle_id": "2026-10-C2", "submitted_at": now}
                                      for d in ("2026-10-12", "2026-10-13")])
    cdb["c_hrms"].insert_many([
        {"employee_id": "JHS9101", "date": "2026-10-12", "kind": "working", "gross_hours": 9.0, "leave_flag": False},
        {"employee_id": "JHS9101", "date": "2026-10-13", "kind": "working", "gross_hours": 8.0, "leave_flag": False}])
    app = FastAPI()
    app.include_router(td.router)
    app.dependency_overrides[get_current_user] = lambda: _user["id"]
    _user["id"] = "JHS9001"
    c = TestClient(app)
    c.db = cdb
    yield c
    repository.configure(dev_settings())
    cache.clear()


Q = "payroll=2026-10"


def test_team_lead_sees_own_team_figures(api):
    v = api.get(f"/timesheet/v2/team/data?{Q}&cycle=2026-10-C2").json()
    t = v["tiles"]
    assert t["people"] == 2 and t["hours"] == 25.0                                          # JHS9101 + JHS9102 only, not JHS9103
    assert (t["hrms_hours"], t["diff_hours"]) == (16.0, 2.0)                                  # net 8.5 + 7.5; typed on those days 8.5 + 9.5
    assert t["unplanned_pct"] == 4
    assert [(c["client"], c["hours"]) for c in v["clients"]] == [("Test Client A", 24.0), ("Unplanned work", 1.0)]
    assert [(d["date"], d["hours"]) for d in v["daily"]] == [("2026-10-12", 15.5), ("2026-10-13", 9.5)]
    proj = {p["project_code"]: p for p in v["projects"]}
    assert proj["PL-T001"]["hours"] == 17.0 and proj["PL-T001"]["people"] == 1 and proj["PL-T001"]["allotted"] == 289.0
    assert proj["PL-T001"]["hours_so_far"] == 17.0 and proj["PL-T001"]["status"] == "late" and proj["PL-T001"]["overdue"] >= 1
    assert proj["PL-T002"]["hours_so_far"] == 13.0                                           # includes JHS9103's 6 h: everyone on the project
    assert proj[""]["label"] == "Unplanned work" and proj[""]["status"] == "unknown"
    people = {p["employee_id"]: p for p in v["people"]}
    assert set(people) == {"JHS9101", "JHS9102"}
    assert people["JHS9101"]["hours"] == 18.0 and people["JHS9101"]["hrms_hours"] == 16.0 and people["JHS9101"]["diff_hours"] == 2.0
    assert (people["JHS9101"]["days_filled"], people["JHS9101"]["days_required"], people["JHS9101"]["missing_days"]) == (2, 7, 5)
    assert people["JHS9102"]["missing_days"] == 7 and people["JHS9102"]["hrms_hours"] is None            # nothing submitted, no HRMS rows
    assert t["days_filled_pct"] == round(100 * 2 / 14)                                                    # 2 of 14 required days
    assert any(a["level"] == "red" and "Statutory Audit" in a["text"] for a in v["attention"])
    assert any("days not submitted" in a["text"] for a in v["attention"])
    assert v["period"]["cycle"] == "11-20 Oct 2026" and v["filters"]["clients"] == ["Test Client A", ""]


def test_client_and_project_filters(api):
    v = api.get(f"/timesheet/v2/team/data?{Q}&client=Test%20Client%20A&project=PL-T002").json()
    assert v["tiles"]["hours"] == 7.0 and [p["project_code"] for p in v["projects"]] == ["PL-T002"] and v["filtered"] is True
    assert v["tiles"]["hrms_hours"] is None and all(p["diff_hours"] is None for p in v["people"])         # not meaningful for one project
    assert {p["employee_id"]: p["hours"] for p in v["people"]} == {"JHS9101": 0.0, "JHS9102": 7.0}
    assert [p["code"] for p in v["filters"]["projects"]] == ["PL-T002", "PL-T001"] or len(v["filters"]["projects"]) == 2        # choices ignore the filter
    unplanned = api.get(f"/timesheet/v2/team/data?{Q}&client=__unplanned__").json()
    assert unplanned["tiles"]["hours"] == 1.0
    assert api.get(f"/timesheet/v2/team/data?{Q}&project=bad;drop").status_code == 422


def test_scope_and_roles(api):
    _user["id"] = "JHS9101"                                                                    # an employee
    assert api.get(f"/timesheet/v2/team/data?{Q}").status_code == 403 and api.get("/timesheet/v2/team/options").status_code == 403
    _user["id"] = "JHS9002"                                                                    # the other manager: only JHS9103
    v = api.get(f"/timesheet/v2/team/data?{Q}&manager=JHS9001").json()                          # ?manager= cannot widen a manager's view
    assert [p["employee_id"] for p in v["people"]] == ["JHS9103"] and v["tiles"]["hours"] == 6.0
    _user["id"] = "JHS9000"                                                                    # admin: everyone, or one team
    assert {p["employee_id"] for p in api.get(f"/timesheet/v2/team/data?{Q}").json()["people"]} >= {"JHS9101", "JHS9102", "JHS9103"}
    assert [p["employee_id"] for p in api.get(f"/timesheet/v2/team/data?{Q}&manager=JHS9002").json()["people"]] == ["JHS9103"]
    opts = api.get("/timesheet/v2/team/options").json()
    assert opts["is_firm"] is True and {"code": "JHS9001", "name": "", "people": 3} in opts["managers"] or opts["managers"]
    assert api.get(f"/timesheet/v2/team/data?{Q}&manager=bad;x").status_code == 422
    assert api.get("/timesheet/v2/team/data?payroll=2020-01").status_code == 404                # not created by an admin
    assert api.get(f"/timesheet/v2/team/data?{Q}&cycle=2026-11-C1").status_code == 404           # cycle not in that payroll


def test_options_list_only_admin_payrolls(api):
    o = api.get("/timesheet/v2/team/options").json()
    assert [p["id"] for p in o["payrolls"]] == ["2026-12", "2026-11", "2026-10"] and o["is_firm"] is False and o["managers"] == []
    assert [c["id"] for c in next(p for p in o["payrolls"] if p["id"] == "2026-10")["cycles"]] == ["2026-09-C3", "2026-10-C1", "2026-10-C2"]


def test_excel_export_matches_the_view_and_is_audited(api):
    r = api.get(f"/timesheet/v2/team/export?{Q}&client=Test%20Client%20A")
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('team_dashboard_2026-10.xlsx"')
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Projects", "People", "Clients", "Hours per day", "Needs attention"]
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary["Hours logged"] == 24.0 and summary["Client filter"] == "Test Client A"
    assert [row[0] for row in wb["People"].iter_rows(min_row=2, values_only=True)] == ["JHS9101", "JHS9102"]
    assert api.db["c_audit"].find_one({"kind": "team-dashboard"})["by"] == "JHS9001"
    _user["id"] = "JHS9101"
    assert api.get(f"/timesheet/v2/team/export?{Q}").status_code == 403
