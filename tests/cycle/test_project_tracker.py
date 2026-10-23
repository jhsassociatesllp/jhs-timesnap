"""Project tracker: the plain-language rules and the API (scope, filters, figures, Excel). Scratch Mongo + the fake NQ in Docker."""
import io
from datetime import date, datetime

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.auth import get_current_user
from backend.metrics import access
from backend.metrics.store import mstore
from backend.nq import cache, repository
from backend.nq.models import NQTask
from backend.timesheet import project_tracker as pt
from tests.cycle.conftest import entry
from tests.nq.conftest import dev_settings

_user = {"id": "JHS9001"}
TODAY = date(2026, 10, 21)


def task(**kw):
    base = dict(id=1, project_id=1, name="Stage", start_date=date(2026, 10, 1), end_date=date(2026, 10, 30), assignee_ids=[3], status_id=1)
    base.update(kw)
    base["is_complete"] = base["status_id"] == 4
    base["is_overdue"] = bool(base["end_date"] and base["end_date"] < TODAY and not base["is_complete"])
    return NQTask(**base)


# ── pure rules ───────────────────────────────────────────────────────────────

def test_audit_type_is_tidied():
    assert pt.tidy_audit_type("  Statutory   Audit ") == "Statutory Audit"
    assert pt.tidy_audit_type("S") == "Not set" and pt.tidy_audit_type(None) == "Not set" and pt.tidy_audit_type("  ") == "Not set"


def test_state_and_live_status():
    done, going, todo, hold = task(status_id=4), task(status_id=2), task(), task(status_id=3)
    assert pt.project_state("active", [done, done]) == "Completed" and pt.project_state("closed", [todo]) == "Completed"
    assert pt.project_state("active", [done, todo]) == "Ongoing" and pt.project_state("active", [going]) == "Ongoing"
    assert pt.project_state("active", [todo, hold]) == "Pending" and pt.project_state("active", []) == "Pending"
    assert pt.project_state("hold", [done]) == "On hold"
    ls = pt.live_status
    assert ls("Completed", "active", date(2026, 1, 1), [done], TODAY) == "completed"
    assert ls("Pending", "active", None, [], TODAY) == "no_plan"
    assert ls("Pending", "hold", date(2026, 12, 1), [todo], TODAY) == "on_hold"
    assert ls("Ongoing", "active", date(2026, 10, 20), [going], TODAY) == "delayed"                   # final report date passed
    assert ls("Ongoing", "active", date(2026, 12, 1), [task(end_date=date(2026, 10, 10))], TODAY) == "delayed"   # an open stage overdue
    assert ls("Ongoing", "active", date(2026, 10, 24), [going], TODAY) == "at_risk"                   # 3 days, earlier stage open
    assert ls("Ongoing", "active", date(2026, 10, 24), [task(name="Final Report", status_id=2)], TODAY) == "on_track"  # only the report is left
    assert ls("Ongoing", "active", date(2026, 11, 20), [going], TODAY) == "on_track"


def test_final_report_deadline():
    fr, other = task(name="Final report issue", end_date=date(2026, 11, 5)), task(end_date=date(2026, 10, 30))
    assert pt.final_report_deadline([other, fr], date(2026, 12, 1)) == date(2026, 11, 5)
    assert pt.final_report_deadline([other], date(2026, 12, 1)) == date(2026, 12, 1)
    assert pt.final_report_deadline([other], None) == date(2026, 10, 30) and pt.final_report_deadline([], None) is None


# ── API ──────────────────────────────────────────────────────────────────────

@pytest.fixture
def api(cdb, monkeypatch):
    repository.configure(dev_settings())
    cache.clear()
    cdb["c_audit"].drop()
    monkeypatch.setattr(mstore, "audit", cdb["c_audit"])
    monkeypatch.setattr(access, "has_admin_module", lambda emp, key: emp == "JHS9000" and key == "timesheet")
    monkeypatch.setattr(pt, "_today", lambda: TODAY)
    cdb["c_entries"].insert_many([
        entry("JHS9101", "2026-10-20", 3, task_id=2, client="Test Client A", project="Statutory Audit FY26", project_code="PL-T001"),
        entry("JHS9102", "2026-10-19", 2, task_id=2, client="Test Client A", project="Statutory Audit FY26", project_code="PL-T001"),
        entry("JHS9102", "2026-10-01", 5, task_id=2, client="Test Client A", project="Statutory Audit FY26", project_code="PL-T001")])   # older than a week
    app = FastAPI()
    app.include_router(pt.router)
    app.dependency_overrides[get_current_user] = lambda: _user["id"]
    _user["id"] = "JHS9001"
    c = TestClient(app)
    c.db = cdb
    yield c
    repository.configure(dev_settings())
    cache.clear()


def by_code(v):
    return {p["code"]: p for p in v["projects"]}


def test_manager_sees_all_his_projects_and_clients(api):
    v = api.get("/timesheet/v2/tracker/data").json()
    k = v["kpis"]
    assert [p["code"] for p in v["projects"]] == ["PL-T002", "PL-T001", "PL-T007", "PL-T003"]       # delayed, delayed, not planned, on track
    assert (k["clients"], k["projects"], k["completed"], k["ongoing"], k["pending"]) == (2, 4, 0, 1, 3)
    assert (k["delayed"], k["at_risk"], k["on_track"]) == (2, 0, 1)
    assert (k["not_planned"], k["partly_planned"], k["not_assigned"], k["partly_assigned"]) == (1, 1, 1, 1)
    assert v["hidden_closed"] == 1 and v["manager"] and "PL-T006" not in by_code(v)                  # closed hidden; another manager's project absent
    p = by_code(v)
    assert p["PL-T001"]["audit_type"] == "Statutory Audit" and p["PL-T001"]["client"] == "Test Client A" and p["PL-T001"]["role"] == "Project manager"
    assert (p["PL-T001"]["state"], p["PL-T001"]["plan"], p["PL-T001"]["assigned"], p["PL-T001"]["status"]) == ("Ongoing", "Planned", "Assigned", "delayed")
    assert (p["PL-T001"]["stages_done"], p["PL-T001"]["stages_total"], p["PL-T001"]["stages_overdue"], p["PL-T001"]["days_left"]) == (1, 4, 2, 10)
    assert p["PL-T001"]["week_hours"] == 5.0 and [w["hours"] for w in p["PL-T001"]["week"]] == [3.0, 2.0]    # last 7 days only
    assert (p["PL-T002"]["assigned"], p["PL-T002"]["state"]) == ("Partly assigned", "Pending")
    assert p["PL-T003"]["plan"] == "Partly planned" and p["PL-T003"]["status"] == "on_track" and p["PL-T003"]["audit_type"] == "Not set"
    assert (p["PL-T007"]["plan"], p["PL-T007"]["assigned"], p["PL-T007"]["status"], p["PL-T007"]["tasks"]) == ("Not planned", "Not assigned", "no_plan", [])
    assert {t["name"]: t["assignees"] for t in p["PL-T001"]["tasks"]}["Planning and scoping"] == ["Test Employee 1"]
    assert [(c["client"], c["projects"], c["delayed"]) for c in v["clients"]] == [("Test Client A", 2, 2), ("Test Client B", 2, 0)]
    assert any("Delayed" in s for s in v["summary"]) and any("No plan yet" in s for s in v["summary"])


def test_show_closed_and_filters(api):
    v = api.get("/timesheet/v2/tracker/data?show=all").json()
    assert v["kpis"]["projects"] == 5 and v["kpis"]["completed"] == 1 and by_code(v)["PL-T005"]["status"] == "completed"
    assert [p["code"] for p in api.get("/timesheet/v2/tracker/data?show=closed").json()["projects"]] == ["PL-T005"]
    assert {p["code"] for p in api.get("/timesheet/v2/tracker/data?client=Test%20Client%20B").json()["projects"]} == {"PL-T003", "PL-T007"}
    assert [p["code"] for p in api.get("/timesheet/v2/tracker/data?audit=Internal%20Audit").json()["projects"]] == ["PL-T002"]
    assert [p["code"] for p in api.get("/timesheet/v2/tracker/data?status=no_plan").json()["projects"]] == ["PL-T007"]
    assert api.get("/timesheet/v2/tracker/data?status=bogus").status_code == 422 and api.get("/timesheet/v2/tracker/data?show=x").status_code == 422
    assert api.get("/timesheet/v2/tracker/data").json()["filters"]["audits"] == ["GST Audit", "Internal Audit", "Not set", "Statutory Audit", "Stock Audit"]


def test_roles_and_scope(api):
    _user["id"] = "JHS9101"                                                                  # an employee
    assert api.get("/timesheet/v2/tracker/data").status_code == 403 and api.get("/timesheet/v2/tracker/export").status_code == 403
    _user["id"] = "JHS9002"                                                                  # another manager (not in NQ); only JHS9103 reports to him
    v = api.get("/timesheet/v2/tracker/data?manager=JHS9001").json()                         # ?manager= cannot widen a manager's view
    assert set(by_code(v)) == {"PL-T002", "PL-T003"} and {p["role"] for p in v["projects"]} == {"Team on tasks"}
    _user["id"] = "JHS9000"                                                                  # admin / partner: the whole firm, or one manager
    v = api.get("/timesheet/v2/tracker/data").json()
    assert v["is_firm"] is True and {"PL-T006", "PL-T001"} <= set(by_code(v)) and v["manager"] == "Whole firm"
    assert set(by_code(api.get("/timesheet/v2/tracker/data?manager=JHS9001").json())) == {"PL-T001", "PL-T002", "PL-T003", "PL-T007"}
    assert api.get("/timesheet/v2/tracker/data?manager=bad;x").status_code == 422


def test_excel_matches_the_view_and_is_audited(api):
    r = api.get("/timesheet/v2/tracker/export?show=all")
    assert r.status_code == 200 and "project_tracker_2026-10-21.xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Dashboard", "Projects", "Master Tracker", "Clients", "Notes"]
    dash = {row[0]: row[1] for row in wb["Dashboard"].iter_rows(min_row=2, values_only=True) if row[0]}
    assert dash["Total projects"] == 5 and dash["Total clients"] == 2 and dash["Completed"] == 1
    assert [row[1] for row in wb["Projects"].iter_rows(min_row=2, values_only=True)][:2] == ["PL-T002", "PL-T001"]
    master = list(wb["Master Tracker"].iter_rows(min_row=2, values_only=True))
    assert any(row[1] == "PL-T007" and "no stages planned" in row[4] for row in master)           # a project with no plan still has a row
    assert any(row[4] == "Overdue open task" and row[8] == "Pending (overdue)" and row[7] == "Test Employee 1" for row in master)
    assert api.db["c_audit"].find_one({"kind": "project-tracker"})["by"] == "JHS9001"
