# backend/timesheet/project_tracker.py
"""
Project tracker for team leads (managers), partners and admins: every project and client linked to a manager, straight
from Nexus Quant (NQ), with a plain status for each one. Calculated when the page is opened.

  GET /timesheet/v2/tracker/data     ?manager=&client=&audit=&status=&show=    everything the page shows
  GET /timesheet/v2/tracker/export   the same view as an Excel file (Dashboard, Projects, Master Tracker, Clients, Notes)

A project belongs to a manager when he is its project manager or partner in NQ, or when he or one of his reportees is
assigned to one of its tasks. Admins and partners see every project, or one manager's with ?manager=.
Per project (all words are kept simple):
  state     Completed (every stage done, or the project is closed) / Ongoing (something started or done) / Pending (nothing
            started) / On hold
  plan      Planned (every stage has dates) / Partly planned / Not planned (no stages or no dates)
  assigned  Assigned (every stage has someone) / Partly assigned / Not assigned
  status    live status: Completed, Delayed, At risk, On track, Not planned, On hold
Hours (who worked this week) come from the new timesheet; the rest is NQ only. Nothing here writes to NQ.
"""
import io
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from typing import Dict, List, Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Font, PatternFill

from backend.auth import get_current_user
from backend.metrics.exports import XLSX, safe_cell
from backend.metrics.store import mstore
from backend.cycle.store import store
from backend.nq import repository as nq_repo
from backend.nq.models import NQTask, NQUnavailable
from backend.timesheet.team_dashboard import _CODE, _today, _who

router = APIRouter(prefix="/timesheet/v2/tracker", tags=["Project tracker"])
AT_RISK_DAYS = 3                       # a final report due within this many days with earlier stages still open
FINAL_REPORT = re.compile(r"final\s*report", re.I)
STATUS_LABELS = {"completed": "Completed", "delayed": "Delayed", "at_risk": "At risk", "on_track": "On track",
                 "no_plan": "Not planned", "on_hold": "On hold"}
SHOW = ("open", "closed", "all")


# ─────────────────────────────── pure rules (tested) ─────────────────────────

def tidy_audit_type(raw: Optional[str]) -> str:
    """NQ category is free text: trim, collapse spaces; a stray one-letter or empty value means not set."""
    s = " ".join(str(raw or "").split())
    return s if len(s) > 1 else "Not set"


def _overdue(t: NQTask, today: date) -> bool:
    return bool(t.end_date and t.end_date < today and not t.is_complete)


def task_state(t: NQTask) -> str:
    if t.is_complete:
        return "Completed"
    if t.status_id == 3:
        return "On hold"
    if t.status_id == 2:
        return "Ongoing"
    return "Pending"


def _level(flags: List[bool], whole: str, some: str, none: str) -> str:
    if flags and all(flags):
        return whole
    return some if any(flags) else none


def project_state(project_status: str, stages: List[NQTask]) -> str:
    if project_status == "closed":
        return "Completed"
    if project_status == "hold":
        return "On hold"
    if stages and all(s.is_complete for s in stages):
        return "Completed"
    if any(s.is_complete or s.status_id == 2 for s in stages):
        return "Ongoing"
    return "Pending"


def live_status(state: str, project_status: str, deadline: Optional[date], stages: List[NQTask], today: date) -> str:
    """'completed' | 'on_hold' | 'no_plan' | 'delayed' | 'at_risk' | 'on_track'."""
    if state == "Completed":
        return "completed"
    if project_status in ("hold", "inactive"):
        return "on_hold"
    if not stages and not deadline:
        return "no_plan"
    open_stages = [s for s in stages if not s.is_complete]
    if (deadline and deadline < today) or any(_overdue(s, today) for s in open_stages):
        return "delayed"
    before_final = [s for s in open_stages if not FINAL_REPORT.search(s.name)]
    if deadline and (deadline - today).days <= AT_RISK_DAYS and before_final:
        return "at_risk"
    return "on_track"


def final_report_deadline(stages: List[NQTask], project_end: Optional[date]) -> Optional[date]:
    ends = [s.end_date for s in stages if s.end_date and FINAL_REPORT.search(s.name)]
    if ends:
        return max(ends)
    if project_end:
        return project_end
    ends = [s.end_date for s in stages if s.end_date]
    return max(ends) if ends else None


# ─────────────────────────────── the view ────────────────────────────────────

def _iso(d: Optional[date]) -> str:
    return d.isoformat() if d else ""


def _associated(me_uid: Optional[int], team_uids: set, projects, tasks_by_pid, all_projects: bool):
    """(project, role) for each project linked to the manager."""
    for p in projects:
        if all_projects:
            yield p, ""
            continue
        if me_uid is not None and me_uid in p.manager_ids:
            yield p, "Project manager"
        elif me_uid is not None and p.partner_id == me_uid:
            yield p, "Partner"
        elif any(set(t.assignee_ids) & team_uids for t in tasks_by_pid.get(p.id, [])):
            yield p, "Team on tasks"


def build_tracker(me: str, firm: bool, reportees: List[str], manager: Optional[str], client: str, audit: str,
                  status: str, show: str) -> dict:
    today = _today()
    out = {"as_of": today.isoformat(), "nq_unavailable": False, "kpis": {}, "summary": [], "clients": [], "projects": [],
           "filters": {"clients": [], "audits": []}, "hidden_closed": 0, "manager": "", "note": ""}
    scope_emp = (manager.strip().upper() if manager else None) if firm else me
    if scope_emp and not _CODE.match(scope_emp):
        raise HTTPException(422, "Manager code: letters, digits, - and _ only")
    team_emps = ([scope_emp] + [d["EmpID"].strip().upper() for d in store.employees.find({"ReportingEmpCode": scope_emp}, {"EmpID": 1}) if d.get("EmpID")]
                 if scope_emp else [])
    if scope_emp:
        d = store.employees.find_one({"EmpID": scope_emp}, {"Emp Name": 1})
        out["manager"] = (d or {}).get("Emp Name") or scope_emp
    else:
        out["manager"] = "Whole firm"

    try:
        users = nq_repo.get_users()
        projects = nq_repo.list_all_projects()
        tasks = nq_repo.list_all_tasks()
    except NQUnavailable:
        out["nq_unavailable"] = True
        return out
    uid_by_emp = {u.emp_id.strip().upper(): u.id for u in users}
    name_by_uid = {u.id: u.name for u in users}
    tasks_by_pid: Dict[int, List[NQTask]] = defaultdict(list)
    for t in tasks:
        tasks_by_pid[t.project_id].append(t)
    me_uid = uid_by_emp.get(scope_emp) if scope_emp else None
    team_uids = {uid_by_emp[e] for e in team_emps if e in uid_by_emp}
    if scope_emp and not team_uids:
        out["note"] = "This manager is not found in Nexus Quant, so no projects can be linked."

    linked = list(_associated(me_uid, team_uids, projects, tasks_by_pid, all_projects=not scope_emp))
    codes = [p.project_code for p, _ in linked]

    # who worked on each project in the last 7 days (new timesheet, everyone)
    week_start = (today - timedelta(days=6)).isoformat()
    week: Dict[str, Counter] = defaultdict(Counter)
    for x in store.entries.aggregate([{"$match": {"project_code": {"$in": codes}, "date": {"$gte": week_start, "$lte": today.isoformat()},
                                                  "hours": {"$ne": None}}},
                                      {"$group": {"_id": {"p": "$project_code", "e": "$employee_id"}, "h": {"$sum": "$hours"}}}]) if codes else []:
        week[x["_id"]["p"]][x["_id"]["e"]] += x["h"]
    emp_names = {d["EmpID"].strip().upper(): d.get("Emp Name") or "" for d in
                 store.employees.find({"EmpID": {"$in": sorted({e for c in week.values() for e in c})}}, {"EmpID": 1, "Emp Name": 1}) if d.get("EmpID")} if week else {}

    rows = []
    for p, role in linked:
        all_tasks = tasks_by_pid.get(p.id, [])
        stages = [t for t in all_tasks if t.parent_task_id is None] or all_tasks
        state = project_state(p.status, stages)
        deadline = final_report_deadline(stages, p.end_date)
        stat = live_status(state, p.status, deadline, stages, today)
        planned = _level([bool(s.start_date and s.end_date) for s in stages], "Planned", "Partly planned", "Not planned")
        assigned = _level([bool(s.assignee_ids) for s in stages], "Assigned", "Partly assigned", "Not assigned")
        open_stages = sorted([s for s in stages if not s.is_complete], key=lambda s: (s.end_date or date.max, s.id))
        task_rows = []
        for t in sorted(all_tasks, key=lambda t: (t.parent_task_id is not None, t.start_date or date.max, t.id)):
            task_rows.append({"name": t.name, "sub_task": t.parent_task_id is not None, "start": _iso(t.start_date), "end": _iso(t.end_date),
                              "assignees": [name_by_uid.get(i, f"NQ user {i}") for i in t.assignee_ids], "state": task_state(t),
                              "planned": bool(t.start_date and t.end_date), "assigned": bool(t.assignee_ids),
                              "overdue": _overdue(t, today)})
        wk = [{"name": emp_names.get(e) or e, "hours": round(h, 1)} for e, h in week.get(p.project_code, Counter()).most_common()]
        rows.append({
            "code": p.project_code, "project": p.project_name or p.project_code, "client": p.client_name or "No client",
            "audit_type": tidy_audit_type(p.category), "role": role, "nq_status": p.status, "state": state, "plan": planned,
            "assigned": assigned, "status": stat, "status_label": STATUS_LABELS[stat],
            "start": _iso(p.start_date), "end": _iso(p.end_date), "deadline": _iso(deadline),
            "days_left": (deadline - today).days if deadline else None,
            "stages_total": len(stages), "stages_done": sum(1 for s in stages if s.is_complete),
            "stages_ongoing": sum(1 for s in stages if not s.is_complete and s.status_id == 2),
            "stages_pending": sum(1 for s in stages if not s.is_complete and s.status_id == 1),
            "stages_hold": sum(1 for s in stages if not s.is_complete and s.status_id == 3),
            "stages_overdue": sum(1 for s in stages if _overdue(s, today)),
            "stages_unassigned": sum(1 for s in stages if not s.assignee_ids),
            "stages_undated": sum(1 for s in stages if not (s.start_date and s.end_date)),
            "next_stage": open_stages[0].name if open_stages else "", "next_stage_due": _iso(open_stages[0].end_date) if open_stages else "",
            "week": wk, "week_hours": round(sum(w["hours"] for w in wk), 1), "tasks": task_rows})

    # filter choices come from everything linked, before the filters below
    closed_state = lambda r: r["nq_status"] == "closed"
    out["filters"] = {"clients": sorted({r["client"] for r in rows}, key=str.lower), "audits": sorted({r["audit_type"] for r in rows}, key=str.lower)}
    if show == "open":
        out["hidden_closed"] = sum(1 for r in rows if closed_state(r))
        rows = [r for r in rows if not closed_state(r)]
    elif show == "closed":
        rows = [r for r in rows if closed_state(r)]
    if client:
        rows = [r for r in rows if r["client"] == client]
    if audit:
        rows = [r for r in rows if r["audit_type"] == audit]
    if status:
        rows = [r for r in rows if r["status"] == status]
    rows.sort(key=lambda r: (["delayed", "at_risk", "no_plan", "on_hold", "on_track", "completed"].index(r["status"]), r["client"].lower(), r["project"].lower()))

    by_client: Dict[str, dict] = {}
    for r in rows:
        c = by_client.setdefault(r["client"], {"client": r["client"], "projects": 0, "completed": 0, "ongoing": 0, "pending": 0,
                                               "on_hold": 0, "not_planned": 0, "not_assigned": 0, "delayed": 0})
        c["projects"] += 1
        c[{"Completed": "completed", "Ongoing": "ongoing", "Pending": "pending", "On hold": "on_hold"}[r["state"]]] += 1
        c["not_planned"] += r["plan"] == "Not planned"
        c["not_assigned"] += r["assigned"] == "Not assigned"
        c["delayed"] += r["status"] == "delayed"
    out["clients"] = sorted(by_client.values(), key=lambda c: (-c["delayed"], -c["projects"], c["client"].lower()))
    cnt = lambda key, val: sum(1 for r in rows if r[key] == val)
    out["kpis"] = {"clients": len(by_client), "projects": len(rows), "completed": cnt("state", "Completed"), "ongoing": cnt("state", "Ongoing"),
                   "pending": cnt("state", "Pending"), "on_hold": cnt("state", "On hold"), "not_planned": cnt("plan", "Not planned"),
                   "partly_planned": cnt("plan", "Partly planned"), "not_assigned": cnt("assigned", "Not assigned"),
                   "partly_assigned": cnt("assigned", "Partly assigned"), "on_track": cnt("status", "on_track"),
                   "at_risk": cnt("status", "at_risk"), "delayed": cnt("status", "delayed"),
                   "stages_total": sum(r["stages_total"] for r in rows), "stages_done": sum(r["stages_done"] for r in rows)}
    out["projects"] = rows
    out["summary"] = _summary(rows, out["kpis"], today)
    return out


def _names(rows, limit=4) -> str:
    names = [f"{r['project']} ({r['client']})" for r in rows[:limit]]
    return ", ".join(names) + (f" and {len(rows) - limit} more" if len(rows) > limit else "")


def _summary(rows: List[dict], k: dict, today: date) -> List[str]:
    if not rows:
        return ["No projects match this view."]
    out = [f"{k['projects']} project{'s' if k['projects'] != 1 else ''} across {k['clients']} client{'s' if k['clients'] != 1 else ''}: "
           f"{k['completed']} completed, {k['ongoing']} ongoing, {k['pending']} pending" + (f", {k['on_hold']} on hold." if k["on_hold"] else ".")]
    for key, text in (("delayed", "Delayed"), ("at_risk", "At risk (final report due within 3 days, earlier stages still open)")):
        hit = [r for r in rows if r["status"] == key]
        if hit:
            out.append(f"{text}: {_names(hit)}.")
    np_ = [r for r in rows if r["plan"] == "Not planned" and r["state"] != "Completed"]
    if np_:
        out.append(f"No plan yet (no stages with dates) for {len(np_)}: {_names(np_)}.")
    na = [r for r in rows if r["assigned"] != "Assigned" and r["state"] != "Completed"]
    if na:
        out.append(f"Stages with nobody assigned in {len(na)} project{'s' if len(na) != 1 else ''}: {_names(na)}.")
    soon = sorted([r for r in rows if r["status"] in ("on_track", "at_risk") and r["deadline"]], key=lambda r: r["deadline"])[:3]
    if soon:
        out.append("Next final reports due: " + "; ".join(f"{r['project']} on {r['deadline']}" for r in soon) + ".")
    if k["delayed"] == 0 and k["at_risk"] == 0:
        out.append("Nothing is delayed or at risk.")
    return out


def _params(manager, client, audit, status, show):
    if show not in SHOW:
        raise HTTPException(422, "show must be open, closed or all")
    if status and status not in STATUS_LABELS:
        raise HTTPException(422, "Unknown status")
    for v in (client, audit):
        if len(v) > 150 or any(ord(ch) < 32 for ch in v):
            raise HTTPException(422, "Bad filter value")
    return manager, client.strip(), audit.strip(), status, show


def _view(current_user, manager, client, audit, status, show):
    me, firm, reportees = _who(current_user)
    manager, client, audit, status, show = _params(manager, client, audit, status, show)
    return me, build_tracker(me, firm, reportees, manager, client, audit, status, show)


# ─────────────────────────────── routes ──────────────────────────────────────

@router.get("/data")
def data(manager: Optional[str] = None, client: str = "", audit: str = "", status: str = "", show: str = "open",
         current_user: str = Depends(get_current_user)):
    me, v = _view(current_user, manager, client, audit, status, show)
    _, firm, _ = _who(current_user)
    v["is_firm"] = firm
    if firm:
        v["managers"] = [{"code": str(d["_id"]).strip().upper(), "name": d.get("name") or "", "people": d["n"]} for d in
                         store.employees.aggregate([{"$match": {"ReportingEmpCode": {"$nin": ["", None]}}},
                                                    {"$group": {"_id": "$ReportingEmpCode", "name": {"$first": "$ReportingEmpName"}, "n": {"$sum": 1}}},
                                                    {"$sort": {"name": 1}}, {"$limit": 500}])]
    return v


@router.get("/export")
def export(manager: Optional[str] = None, client: str = "", audit: str = "", status: str = "", show: str = "open",
           current_user: str = Depends(get_current_user)):
    me, v = _view(current_user, manager, client, audit, status, show)
    if v["nq_unavailable"]:
        raise HTTPException(503, "Project data from Nexus Quant is unavailable right now")
    wb = openpyxl.Workbook()
    head_fill, head_font = PatternFill("solid", fgColor="2C3E50"), Font(bold=True, color="FFFFFF")

    def sheet(title, header, rows, first=False, widths=None):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font, c.alignment = head_fill, head_font, Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append([safe_cell(x) for x in r])
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    k = v["kpis"]
    ws = sheet("Dashboard", ["Item", "Value"], [
        ["Manager", v["manager"]], ["As of", v["as_of"]], ["Total clients", k["clients"]], ["Total projects", k["projects"]],
        ["Completed", k["completed"]], ["Ongoing", k["ongoing"]], ["Pending", k["pending"]], ["On hold", k["on_hold"]],
        ["Not planned", k["not_planned"]], ["Partly planned", k["partly_planned"]], ["Not assigned", k["not_assigned"]],
        ["Partly assigned", k["partly_assigned"]], ["On track", k["on_track"]], ["At risk", k["at_risk"]], ["Delayed", k["delayed"]],
        ["Stages done", f"{k['stages_done']} of {k['stages_total']}"], [], ["Management summary", ""]] + [["", s] for s in v["summary"]],
        first=True, widths=[24, 110])
    ws["A1"].fill = head_fill
    sheet("Projects", ["Client", "Project code", "Project", "Audit type", "Your link", "State", "Plan", "Assigned", "Status",
                       "Stages done", "Stages total", "Ongoing", "Pending", "Overdue", "Final report due", "Days left", "Next stage",
                       "Worked this week (hours)"],
          [[p["client"], p["code"], p["project"], p["audit_type"], p["role"], p["state"], p["plan"], p["assigned"], p["status_label"],
            p["stages_done"], p["stages_total"], p["stages_ongoing"], p["stages_pending"], p["stages_overdue"], p["deadline"], p["days_left"],
            p["next_stage"], ", ".join(f"{w['name']} {w['hours']}" for w in p["week"])] for p in v["projects"]],
          widths=[28, 12, 34, 22, 16, 11, 14, 15, 12, 9, 9, 9, 9, 9, 14, 9, 30, 40])
    master = []
    for p in v["projects"]:
        if not p["tasks"]:
            master.append([p["client"], p["code"], p["project"], p["audit_type"], "— no stages planned —", "", "", "", "Pending", "Not planned", "Not assigned", p["status_label"]])
        for t in p["tasks"]:
            master.append([p["client"], p["code"], p["project"], p["audit_type"], ("   ↳ " if t["sub_task"] else "") + t["name"], t["start"], t["end"],
                           ", ".join(t["assignees"]), t["state"] + (" (overdue)" if t["overdue"] else ""),
                           "Planned" if t["planned"] else "Not planned", "Assigned" if t["assigned"] else "Not assigned", p["status_label"]])
    sheet("Master Tracker", ["Client", "Project code", "Project", "Audit type", "Stage / task", "Start", "End", "Assigned to", "State", "Plan",
                             "Assigned", "Project status"], master, widths=[28, 12, 34, 22, 40, 12, 12, 34, 18, 13, 14, 13])
    sheet("Clients", ["Client", "Projects", "Completed", "Ongoing", "Pending", "On hold", "Not planned", "Not assigned", "Delayed"],
          [[c["client"], c["projects"], c["completed"], c["ongoing"], c["pending"], c["on_hold"], c["not_planned"], c["not_assigned"], c["delayed"]]
           for c in v["clients"]], widths=[34, 10, 11, 10, 10, 10, 12, 13, 10])
    sheet("Notes", ["Term", "Meaning"], [
        ["Source", "Nexus Quant (clients, projects, stages, dates, assignees); hours worked this week from the Timesnap timesheet."],
        ["Your link", "Project manager or Partner in NQ, or 'Team on tasks' when he or his reportees are assigned to its stages."],
        ["Stage", "A parent task of the project in NQ. Sub-tasks are listed under their stage in the Master Tracker."],
        ["State", "Completed: every stage done (or the project is closed). Ongoing: some stage started or done. Pending: nothing started. On hold: project on hold."],
        ["Plan", "Planned: every stage has start and end dates. Partly planned: some do. Not planned: none, or no stages."],
        ["Assigned", "Assigned: every stage has someone. Partly assigned: some do. Not assigned: none, or no stages."],
        ["Final report due", "End date of the stage named 'Final Report'; if there is none, the project end date."],
        ["Delayed", "Final report date passed, or an open stage is past its end date."],
        ["At risk", f"Final report due within {AT_RISK_DAYS} days while stages before it are still open."],
        ["On track", "Everything else that is open."],
        ["Audit type", "Taken from the project's category in NQ."]], widths=[20, 120])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    mstore.audit.insert_one({"kind": "project-tracker", "by": me, "manager": manager or "", "client": client, "audit": audit, "status": status,
                             "show": show, "rows": len(v["projects"]), "at": datetime.utcnow()})
    return StreamingResponse(buf, media_type=XLSX, headers={
        "Content-Disposition": f'attachment; filename="project_tracker_{_today().isoformat()}.xlsx"', "Cache-Control": "no-store"})
