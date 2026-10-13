# backend/resume_ranker/extract.py
"""
Text extraction for PDF / DOCX / DOC resumes.

The file's *content* decides how it is read, not its extension: ".doc" files in
the wild are frequently really DOCX, RTF or HTML, and a ".pdf" is sometimes a
Word file that was renamed. Every failure raises ExtractionError with a message
HR can act on — the caller records it against that one file and moves on.
"""
from __future__ import annotations

import io
import os
import re
import shutil
import struct
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET

from backend.resume_ranker.config import MAX_TEXT_CHARS

MAX_PDF_PAGES = 30
_MIN_USEFUL_CHARS = 40

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


class ExtractionError(Exception):
    """The file could not be turned into usable text."""


# ── public entry point ──────────────────────────────────────────────────────

def extract_text(data: bytes, filename: str = "") -> str:
    if not data:
        raise ExtractionError("File is empty")

    kind = sniff_kind(data)
    try:
        if kind == "pdf":
            text = _pdf_text(data)
        elif kind == "docx":
            text = _docx_text(data)
        elif kind == "doc":
            text = _doc_text(data)
        elif kind == "rtf":
            text = _rtf_text(data)
        elif kind == "html":
            text = _html_text(data)
        else:
            raise ExtractionError(
                "Unrecognised file content — not a PDF, DOCX or DOC file (it may be corrupt)"
            )
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"Could not read {kind.upper()} file: {exc}") from exc

    text = clean_text(text)
    if len(text) < _MIN_USEFUL_CHARS:
        hint = " It looks like a scanned image — OCR is not supported." if kind == "pdf" else ""
        raise ExtractionError("No readable text found in the file." + hint)
    return text


def sniff_kind(data: bytes) -> str:
    head = data[:8]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK"):
        return "docx" if _is_docx(data) else "unknown"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "doc"
    if head.startswith(b"{\\rtf"):
        return "rtf"
    if data[:512].lstrip().lower().startswith((b"<html", b"<!doctype html", b"<?xml")):
        return "html"
    # %PDF is allowed to appear a little way into the file
    if b"%PDF" in data[:1024]:
        return "pdf"
    return "unknown"


def _is_docx(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return "word/document.xml" in z.namelist()
    except zipfile.BadZipFile:
        return False


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ").replace("\u2022", "•")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_TEXT_CHARS]


# ── PDF ─────────────────────────────────────────────────────────────────────

def _pdf_text(data: bytes) -> str:
    text = ""
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if getattr(pdf, "doc", None) is not None and getattr(pdf.doc, "encryption", None):
                raise ExtractionError("PDF is password-protected")
            pages = [(p.extract_text() or "") for p in pdf.pages[:MAX_PDF_PAGES]]
        text = "\n\n".join(pages)
    except ExtractionError:
        raise
    except Exception:
        text = ""

    if len(text.strip()) < _MIN_USEFUL_CHARS:  # pdfplumber failed or came back empty — try pypdf
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                try:
                    if not reader.decrypt(""):
                        raise ExtractionError("PDF is password-protected")
                except ExtractionError:
                    raise
                except Exception:
                    raise ExtractionError("PDF is password-protected")
            text = "\n\n".join((p.extract_text() or "") for p in reader.pages[:MAX_PDF_PAGES])
        except ExtractionError:
            raise
        except Exception as exc:
            if not text.strip():
                raise ExtractionError(f"PDF is corrupt or unreadable ({exc})") from exc
    return text


# ── DOCX ────────────────────────────────────────────────────────────────────

def _docx_text(data: bytes) -> str:
    """Walk word/document.xml directly. Unlike python-docx's paragraph list this
    also picks up text boxes and nested tables (common in designed CVs), and it
    skips mc:Fallback so text-box content isn't emitted twice."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            parts = ["word/document.xml"]
            parts += sorted(n for n in z.namelist() if re.match(r"word/(header|footer)\d*\.xml$", n))
            chunks = []
            for name in parts:
                chunks.append(_wordml_to_text(z.read(name)))
        return "\n".join(c for c in chunks if c)
    except Exception:
        return _docx_via_python_docx(data)


def _wordml_to_text(xml_bytes: bytes) -> str:
    root = ET.fromstring(xml_bytes)
    lines: list[str] = []

    def walk(el, buf):
        for child in el:
            tag = child.tag
            if tag == _MC_FALLBACK:
                continue
            if tag == _W + "p":
                para: list[str] = []
                walk_par(child, para)
                line = "".join(para).strip()
                if line:
                    lines.append(line)
                # text boxes hold their own paragraphs — collected by walk_par's recursion
            else:
                walk(child, buf)

    def walk_par(el, out):
        for child in el:
            tag = child.tag
            if tag == _MC_FALLBACK:
                continue
            if tag == _W + "t":
                out.append(child.text or "")
            elif tag == _W + "tab":
                out.append("\t")
            elif tag in (_W + "br", _W + "cr"):
                out.append("\n")
            elif tag == _W + "txbxContent":
                walk(child, None)
            else:
                walk_par(child, out)

    walk(root, None)
    return "\n".join(lines)


def _docx_via_python_docx(data: bytes) -> str:
    import docx
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


# ── legacy DOC ──────────────────────────────────────────────────────────────

def _doc_text(data: bytes) -> str:
    text = _doc_via_antiword(data)
    if text and len(text.strip()) >= _MIN_USEFUL_CHARS:
        return text
    try:
        text = _doc_via_piece_table(data)
        if text and len(text.strip()) >= _MIN_USEFUL_CHARS:
            return text
    except Exception:
        pass
    text = _doc_via_strings(data)
    if len(text.strip()) < _MIN_USEFUL_CHARS:
        raise ExtractionError("Could not read text from this legacy .doc file — try re-saving it as DOCX or PDF")
    return text


def _doc_via_antiword(data: bytes) -> str:
    exe = shutil.which("antiword")
    if not exe:
        return ""
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
        tmp.write(data)
        path = tmp.name
    try:
        proc = subprocess.run([exe, "-m", "UTF-8.txt", path], capture_output=True, timeout=30)
        return proc.stdout.decode("utf-8", "replace") if proc.returncode == 0 else ""
    except Exception:
        return ""
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


def _doc_via_piece_table(data: bytes) -> str:
    """Pure-Python Word 97-2003 reader: follows the FIB to the CLX piece table
    and stitches the text pieces together ([MS-DOC] 2.4.1)."""
    import olefile
    ole = olefile.OleFileIO(io.BytesIO(data))
    try:
        word = ole.openstream("WordDocument").read()
        flags = struct.unpack_from("<H", word, 0x0A)[0]
        table_name = "1Table" if flags & 0x0200 else "0Table"
        table = ole.openstream(table_name).read()
        fc_clx, lcb_clx = struct.unpack_from("<II", word, 0x01A2)
    finally:
        ole.close()

    clx = table[fc_clx:fc_clx + lcb_clx]
    pos = 0
    while pos < len(clx) and clx[pos] == 0x01:            # skip Prc entries
        pos += 3 + struct.unpack_from("<H", clx, pos + 1)[0]
    if pos >= len(clx) or clx[pos] != 0x02:
        raise ValueError("no piece table")
    lcb = struct.unpack_from("<I", clx, pos + 1)[0]
    plc = clx[pos + 5:pos + 5 + lcb]
    n = (lcb - 4) // 12                                   # (n+1) CPs of 4 bytes + n PCDs of 8 bytes
    cps = struct.unpack_from(f"<{n + 1}I", plc, 0)
    out = []
    for i in range(n):
        fc = struct.unpack_from("<I", plc, (n + 1) * 4 + i * 8 + 2)[0]
        length = cps[i + 1] - cps[i]
        if fc & 0x40000000:                               # compressed: 1 byte per char (cp1252)
            start = (fc & 0x3FFFFFFF) // 2
            out.append(word[start:start + length].decode("cp1252", "replace"))
        else:                                             # UTF-16LE
            out.append(word[fc:fc + length * 2].decode("utf-16-le", "replace"))
    text = "".join(out)
    # Word control characters: \r paragraph, \x07 cell/row end, \x0b line break, \x0c page break
    text = text.replace("\x07", " | ").replace("\r", "\n").replace("\x0b", "\n").replace("\x0c", "\n\n")
    text = re.sub(r"[\x01-\x08\x0e-\x1f]", "", text)
    return text


def _doc_via_strings(data: bytes) -> str:
    """Last resort: pull UTF-16LE and 8-bit printable runs out of the raw bytes."""
    runs = re.findall(rb"(?:[\x20-\x7e]\x00){6,}", data)
    text = "\n".join(r.decode("utf-16-le", "ignore") for r in runs)
    if len(text) < _MIN_USEFUL_CHARS:
        runs = re.findall(rb"[\x20-\x7e]{12,}", data)
        text = "\n".join(r.decode("ascii", "ignore") for r in runs)
    return text


# ── disguised formats ───────────────────────────────────────────────────────

def _rtf_text(data: bytes) -> str:
    raw = data.decode("cp1252", "replace")
    try:
        from striprtf.striprtf import rtf_to_text
        return rtf_to_text(raw)
    except Exception:
        return re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", raw)


def _html_text(data: bytes) -> str:
    raw = data.decode("utf-8", "replace")
    raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>|</(p|div|li|tr|h\d)>", "\n", raw)
    raw = re.sub(r"<[^>]+>", " ", raw)
    import html
    return html.unescape(raw)
