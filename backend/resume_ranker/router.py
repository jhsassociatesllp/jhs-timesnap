# backend/resume_ranker/router.py
"""
Resume Ranker API  (prefix: /resume-ranker)

  access-check / overview
  jobs        CRUD, criteria extraction from a pasted JD, (re)scoring, stage changes, export
  batches     bulk upload session + progress polling + retry of failed files
  resumes     search / filter / ranked lists, detail, download, tags, notes, delete

All endpoints are plain `def` (FastAPI runs them in its thread pool), so the
blocking pymongo / file / openpyxl calls never stall the event loop.
"""
from __future__ import annotations

import csv
import hashlib
import io
import logging
import mimetypes
import os
import re
import secrets
from datetime import datetime, timedelta
from typing import Dict, List, Optional

from bson import ObjectId
from bson.errors import InvalidId
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, Field, field_validator

from concurrent.futures import ThreadPoolExecutor

from backend.auth import get_current_user
from backend.database import employee_details_collection
from backend.resume_ranker import config, llm, worker
from backend.resume_ranker.db import (OWNER_EMPID, adhoc_scores, batches, ensure_indexes, holders, jobs, resumes,
                                      scores, tasks)
from backend.resume_ranker.embeddings import cosine_similarity, embed_job
from backend.resume_ranker.parser import EDU_LABELS, EDU_RANKS, general_score, skill_key
from backend.resume_ranker.prefilter import screened_out_result, split_candidates
from backend.resume_ranker.scoring import (CRITERIA, CRITERION_LABELS, DEFAULT_WEIGHTS, active_criteria,
                                           criteria_signature, effective_weights, extract_criteria, score_resume)

log = logging.getLogger("resume_ranker.router")

router = APIRouter(prefix="/resume-ranker", tags=["Resume Ranker"])

STAGES = ("new", "shortlisted", "interviewed", "rejected", "hired")
LIST_EXCLUDE = {"text": 0, "parsed.achievements": 0, "embedding": 0}
# Only high-confidence duplicates (same email / phone / identical file) are hidden and left unscored.
# A "possible" duplicate (same name + shared employer) is just badged — hiding a genuinely
# different person who happens to share a name and employer would silently drop a candidate.
NOT_HARD_DUP = {"$or": [{"duplicate_of": {"$exists": False}}, {"duplicate_confidence": "possible"}]}


# ─────────────────────────────────────────────────────────────────────────────
# Startup + access control
# ─────────────────────────────────────────────────────────────────────────────

def backfill_general_scores() -> int:
    """CVs parsed before automatic ranking existed have no rank score. Compute it from the
    already-stored parsed data (plain arithmetic: no AI call, no cost). Returns how many were filled."""
    n = 0
    for r in resumes.find({"status": "parsed", "general_score": {"$exists": False}, "parsed": {"$exists": True}},
                          {"parsed": 1}):
        gs = general_score(r["parsed"])
        resumes.update_one({"_id": r["_id"]}, {"$set": {"general_score": gs["total"], "general_breakdown": gs["breakdown"]}})
        n += 1
    return n


def startup() -> None:
    """Called once from main.py's startup hook."""
    ensure_indexes()
    try:
        filled = backfill_general_scores()
        if filled:
            log.info("ranked %d previously-uploaded CV(s) (free, no AI)", filled)
    except Exception:
        log.exception("general-score backfill skipped")
    os.makedirs(config.upload_dir(), exist_ok=True)
    worker.pool.start()


def _has_access(empid: str) -> bool:
    """Only people listed in the Resume_Ranker.Holders collection may use the module."""
    empid = empid.strip().upper()
    if holders.find_one({"empid": empid}):
        return True
    # Tolerate hand-added docs: any letter case, stray spaces, an "EmpID" field name,
    # or one document holding an array of codes: {"empids": ["JHS1", "JHS2"]}.
    rx = {"$regex": rf"^\s*{re.escape(empid)}\s*$", "$options": "i"}
    return holders.find_one({"$or": [{"empid": rx}, {"EmpID": rx}, {"empids": rx}]}) is not None


def require_access(current_user: str = Depends(get_current_user)) -> str:
    if not _has_access(current_user):
        raise HTTPException(403, "You do not have access to the Resume Ranker module")
    return current_user.strip().upper()


@router.get("/access-check")
def access_check(current_user: str = Depends(get_current_user)):
    return {"allowed": _has_access(current_user), "ai_configured": llm.is_configured()}


def _require_owner(current_user: str = Depends(get_current_user)) -> str:
    if current_user.strip().upper() != OWNER_EMPID:
        raise HTTPException(403, "Only the owner can change the holders list")
    return current_user.strip().upper()


class HolderIn(BaseModel):
    empid: str
    role: str = "hr"


@router.get("/holders")
def list_holders(_: str = Depends(require_access)):
    return {"holders": [{k: v for k, v in h.items() if k != "_id"} for h in holders.find({}).sort("empid", 1)]}


@router.post("/holders")
def add_holder(body: HolderIn, user: str = Depends(_require_owner)):
    empid = body.empid.strip().upper()
    emp = employee_details_collection.find_one({"EmpID": empid})
    if not emp:
        raise HTTPException(400, "Employee does not exist")
    holders.update_one({"empid": empid}, {"$setOnInsert": {
        "empid": empid, "name": (emp.get("Emp Name") or "").strip(), "role": body.role,
        "designation": emp.get("Designation Name"), "added_by": user, "added_at": datetime.utcnow()}}, upsert=True)
    return {"success": True, "empid": empid}


@router.delete("/holders/{empid}")
def remove_holder(empid: str, _: str = Depends(_require_owner)):
    empid = empid.strip().upper()
    if empid == OWNER_EMPID:
        raise HTTPException(400, "The owner cannot be removed")
    holders.delete_one({"empid": empid})
    return {"success": True}


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _oid(value: str, what: str = "id") -> ObjectId:
    try:
        return ObjectId(value)
    except (InvalidId, TypeError):
        raise HTTPException(404, f"{what} not found")


def _iso(dt) -> Optional[str]:
    return dt.isoformat() + "Z" if isinstance(dt, datetime) else None


def _get_job(job_id: str) -> dict:
    job = jobs.find_one({"_id": _oid(job_id, "Job")})
    if not job:
        raise HTTPException(404, "Job not found")
    return job


def _get_resume(resume_id: str, projection=None) -> dict:
    doc = resumes.find_one({"_id": _oid(resume_id, "Resume")}, projection)
    if not doc:
        raise HTTPException(404, "Resume not found")
    return doc


def _require_ai() -> None:
    if not llm.is_configured():
        raise HTTPException(503, "AI is not configured on the server: set RESUME_OPENAI_API_KEY in .env")


def _clean_list(items, limit=40) -> List[str]:
    seen, out = set(), []
    for it in items or []:
        s = re.sub(r"\s+", " ", str(it)).strip()
        if s and s.lower() not in seen:
            seen.add(s.lower())
            out.append(s[:120])
    return out[:limit]


def _job_view(job: dict, with_counts: bool = True) -> dict:
    # "embedding" (a ~1536-float vector) is internal to the pre-filter and never useful to the
    # frontend — leaving it in would bloat every job response for no reason.
    out = {k: v for k, v in job.items() if k not in ("_id", "embedding")}
    out["id"] = str(job["_id"])
    out["created_at"], out["updated_at"] = _iso(job.get("created_at")), _iso(job.get("updated_at"))
    out["active_criteria"] = active_criteria(job)
    out["effective_weights"] = effective_weights(job)
    out["prefilter_enabled"] = bool(job.get("embedding"))
    if with_counts:
        jid = str(job["_id"])
        scored = scores.count_documents({"job_id": jid, "status": "scored"})
        top = scores.find_one({"job_id": jid, "status": "scored"}, sort=[("total", -1)])
        out["scored_count"] = scored
        out["screened_out_count"] = scores.count_documents({"job_id": jid, "status": "screened_out"})
        out["stale_count"] = scores.count_documents(
            {"job_id": jid, "status": "scored", "criteria_sig": {"$ne": job.get("criteria_sig")}})
        out["shortlisted_count"] = scores.count_documents({"job_id": jid, "stage": "shortlisted"})
        out["top_score"] = top["total"] if top else None
    return out


# ─────────────────────────────────────────────────────────────────────────────
# Jobs
# ─────────────────────────────────────────────────────────────────────────────

class Requirement(BaseModel):
    name: str
    must_have: bool = False


class JobIn(BaseModel):
    title: str = Field(min_length=2, max_length=150)
    department: Optional[str] = None
    description: str = ""
    required_skills: List[Requirement] = []
    preferred_skills: List[str] = []
    min_years: float = Field(default=0, ge=0, le=60)
    max_years: Optional[float] = Field(default=None, ge=0, le=60)
    experience_must_have: bool = False
    education_level: str = "none"
    education_fields: List[str] = []
    education_must_have: bool = False
    required_certifications: List[Requirement] = []
    preferred_certifications: List[str] = []
    industries: List[str] = []
    location: Optional[str] = None
    location_type: str = "onsite"
    location_must_have: bool = False
    weights: Dict[str, float] = Field(default_factory=lambda: dict(DEFAULT_WEIGHTS))
    mask_personal_info: bool = False
    must_have_action: str = "flag"
    must_have_cap: float = Field(default=50, ge=0, le=100)
    status: str = "open"
    score_existing: bool = False        # create/update: queue scoring of resumes already in the database

    @field_validator("education_level")
    @classmethod
    def _edu(cls, v):
        if v not in EDU_RANKS:
            raise ValueError(f"education_level must be one of {sorted(EDU_RANKS)}")
        return v

    @field_validator("location_type")
    @classmethod
    def _loc(cls, v):
        if v not in ("onsite", "hybrid", "remote"):
            raise ValueError("location_type must be onsite, hybrid or remote")
        return v

    @field_validator("must_have_action")
    @classmethod
    def _act(cls, v):
        if v not in ("flag", "cap"):
            raise ValueError("must_have_action must be flag or cap")
        return v

    @field_validator("status")
    @classmethod
    def _status(cls, v):
        if v not in ("open", "closed"):
            raise ValueError("status must be open or closed")
        return v

    @field_validator("weights")
    @classmethod
    def _weights(cls, v):
        bad = set(v) - set(CRITERIA)
        if bad:
            raise ValueError(f"unknown weight keys: {sorted(bad)}")
        if any(x < 0 for x in v.values()):
            raise ValueError("weights cannot be negative")
        return v


def _job_doc(body: JobIn) -> dict:
    def reqs(items):
        seen, out = set(), []
        for r in items:
            name = re.sub(r"\s+", " ", r.name).strip()
            if name and name.lower() not in seen:
                seen.add(name.lower())
                out.append({"name": name[:120], "must_have": bool(r.must_have)})
        return out

    doc = {
        "title": body.title.strip(), "department": (body.department or "").strip() or None,
        "description": body.description.strip()[:20000],
        "required_skills": reqs(body.required_skills), "preferred_skills": _clean_list(body.preferred_skills),
        "min_years": body.min_years, "max_years": body.max_years or None,
        "experience_must_have": body.experience_must_have,
        "education_level": body.education_level, "education_fields": _clean_list(body.education_fields, 10),
        "education_must_have": body.education_must_have,
        "required_certifications": reqs(body.required_certifications),
        "preferred_certifications": _clean_list(body.preferred_certifications, 20),
        "industries": _clean_list(body.industries, 10),
        "location": (body.location or "").strip() or None, "location_type": body.location_type,
        "location_must_have": body.location_must_have,
        "weights": {**DEFAULT_WEIGHTS, **body.weights},
        "mask_personal_info": body.mask_personal_info,
        "must_have_action": body.must_have_action, "must_have_cap": body.must_have_cap,
        "status": body.status,
    }
    if not active_criteria(doc):
        raise HTTPException(422, "Add at least one requirement (skills, experience, education, certifications, "
                                 "industry or a location) — otherwise there is nothing to score against.")
    doc["criteria_sig"] = criteria_signature(doc)
    # Best-effort, powers the relevance pre-filter only (see _queue_scoring) —
    # a failure here never blocks saving the job, it just means this job can't
    # be pre-filtered and every candidate will always be fully LLM-scored.
    try:
        doc["embedding"] = embed_job(doc)
    except Exception:
        log.exception("job embedding failed for %r", doc.get("title"))
    return doc


@router.get("/jobs")
def list_jobs(status: Optional[str] = None, _: str = Depends(require_access)):
    flt = {"status": status} if status in ("open", "closed") else {}
    return {"jobs": [_job_view(j) for j in jobs.find(flt).sort("created_at", -1)]}


@router.post("/jobs")
def create_job(body: JobIn, user: str = Depends(require_access)):
    doc = _job_doc(body)
    doc.update({"created_by": user, "created_at": datetime.utcnow(), "updated_at": datetime.utcnow()})
    jid = jobs.insert_one(doc).inserted_id
    result = _queue_scoring(str(jid), "unscored") if body.score_existing else {"queued": 0, "screened_out": 0}
    return {"job": _job_view(jobs.find_one({"_id": jid})), **result}


@router.get("/jobs/{job_id}")
def get_job(job_id: str, _: str = Depends(require_access)):
    return {"job": _job_view(_get_job(job_id))}


@router.put("/jobs/{job_id}")
def update_job(job_id: str, body: JobIn, user: str = Depends(require_access)):
    old = _get_job(job_id)
    doc = _job_doc(body)
    doc["updated_at"] = datetime.utcnow()
    jobs.update_one({"_id": old["_id"]}, {"$set": doc})
    changed = doc["criteria_sig"] != old.get("criteria_sig")
    result = _queue_scoring(job_id, "all") if (changed and body.score_existing) else {"queued": 0, "screened_out": 0}
    return {"job": _job_view(jobs.find_one({"_id": old["_id"]})), "criteria_changed": changed, **result}


@router.delete("/jobs/{job_id}")
def delete_job(job_id: str, _: str = Depends(require_access)):
    job = _get_job(job_id)
    scores.delete_many({"job_id": job_id})
    tasks.delete_many({"kind": "score", "job_id": job_id, "status": "queued"})
    jobs.delete_one({"_id": job["_id"]})
    return {"success": True}


class ExtractIn(BaseModel):
    description: str = Field(min_length=30, max_length=20000)
    title: Optional[str] = None


@router.post("/jobs/extract-criteria")
def extract_job_criteria(body: ExtractIn, _: str = Depends(require_access)):
    """Paste a job description -> suggested structured criteria to review and edit."""
    _require_ai()
    try:
        return {"criteria": extract_criteria(body.description, body.title)}
    except llm.LLMError as exc:
        raise HTTPException(502, str(exc))
    except llm.LLMTransientError as exc:
        raise HTTPException(503, f"AI is busy, try again in a moment ({exc})")


class ScoreIn(BaseModel):
    scope: str = "unscored"                 # unscored | stale | all | selected
    resume_ids: List[str] = []


def _queue_scoring(job_id: str, scope: str, resume_ids: Optional[List[str]] = None) -> dict:
    """Returns {"queued": n, "screened_out": m}.

    scope="selected" is HR explicitly picking specific candidates (e.g. the
    "Score anyway" override on an already screened-out one) — it ALWAYS fully
    scores, the pre-filter never applies there. Every other scope scores
    against the job's whole eligible pool, which is exactly the case the
    pre-filter exists for (see config.py)."""
    proj = {"_id": 1, "embedding": 1} if scope != "selected" else {"_id": 1}
    base = {"status": "parsed"}
    if scope == "selected":
        base["_id"] = {"$in": [_oid(r, "Resume") for r in resume_ids or []]}
    else:
        base.update(NOT_HARD_DUP)                            # score each candidate once
    pool = list(resumes.find(base, proj))
    if scope in ("unscored", "stale"):
        good = {"job_id": job_id, "status": {"$in": ["scored", "screened_out"]}}
        if scope == "stale":
            good["criteria_sig"] = jobs.find_one({"_id": ObjectId(job_id)}).get("criteria_sig")
        done = {s["resume_id"] for s in scores.find(good, {"resume_id": 1})}
        pool = [r for r in pool if str(r["_id"]) not in done]

    if scope == "selected":
        to_score, screened = pool, []
    else:
        job = jobs.find_one({"_id": ObjectId(job_id)})
        to_score, screened = split_candidates(job, pool)
        now = datetime.utcnow()
        for r in screened:
            scores.update_one(
                {"job_id": job_id, "resume_id": str(r["_id"])},
                {"$set": {**screened_out_result(job, r), "status": "screened_out", "error": None,
                         "criteria_sig": job.get("criteria_sig"), "scored_at": now},
                 "$setOnInsert": {"stage": "new", "created_at": now}},
                upsert=True)
    for r in to_score:
        worker.enqueue("score", r["_id"], job_id)
    return {"queued": len(to_score), "screened_out": len(screened)}


@router.post("/jobs/{job_id}/score")
def score_job(job_id: str, body: ScoreIn, _: str = Depends(require_access)):
    _get_job(job_id)
    _require_ai()
    if body.scope not in ("unscored", "stale", "all", "selected"):
        raise HTTPException(422, "scope must be unscored, stale, all or selected")
    return _queue_scoring(job_id, body.scope, body.resume_ids)


class StageIn(BaseModel):
    resume_ids: List[str] = Field(min_length=1)
    stage: str

    @field_validator("stage")
    @classmethod
    def _stage(cls, v):
        if v not in STAGES:
            raise ValueError(f"stage must be one of {STAGES}")
        return v


@router.put("/jobs/{job_id}/stage")
def set_stage(job_id: str, body: StageIn, _: str = Depends(require_access)):
    _get_job(job_id)
    res = scores.update_many({"job_id": job_id, "resume_id": {"$in": body.resume_ids}}, {"$set": {"stage": body.stage}})
    return {"updated": res.modified_count}


# ─────────────────────────────────────────────────────────────────────────────
# Bulk upload
# ─────────────────────────────────────────────────────────────────────────────

class BatchIn(BaseModel):
    job_ids: List[str] = []
    expected_files: int = Field(default=0, ge=0, le=5000)


@router.post("/batches")
def create_batch(body: BatchIn, user: str = Depends(require_access)):
    _require_ai()
    for jid in body.job_ids:
        _get_job(jid)
    batch_id = secrets.token_hex(8)
    batches.insert_one({"batch_id": batch_id, "job_ids": body.job_ids, "expected_files": body.expected_files,
                        "created_by": user, "created_at": datetime.utcnow(), "rejected": []})
    return {"batch_id": batch_id}


@router.post("/batches/{batch_id}/files")
def upload_files(batch_id: str, files: List[UploadFile] = File(...), user: str = Depends(require_access)):
    """Accepts a handful of files per request (the browser sends a big selection
    in several requests). Each file is validated on its own: one bad file is
    reported and skipped, never fails the others."""
    batch = batches.find_one({"batch_id": batch_id})
    if not batch:
        raise HTTPException(404, "Upload batch not found")
    if len(files) > config.MAX_FILES_PER_REQUEST:
        raise HTTPException(413, f"Send at most {config.MAX_FILES_PER_REQUEST} files per request")

    results, rejected = [], []
    for f in files:
        name = os.path.basename(f.filename or "unnamed")
        ext = os.path.splitext(name)[1].lower()
        reason = None
        data = b""
        if ext not in config.ALLOWED_EXTENSIONS:
            reason = "Unsupported file type (use PDF, DOCX or DOC)"
        else:
            data = f.file.read(config.MAX_FILE_BYTES + 1)
            if not data:
                reason = "File is empty"
            elif len(data) > config.MAX_FILE_BYTES:
                reason = f"File is larger than {config.MAX_FILE_BYTES // (1024 * 1024)} MB"
        if reason:
            rejected.append({"filename": name, "reason": reason})
            results.append({"filename": name, "status": "rejected", "reason": reason})
            continue
        try:
            results.append(_store_resume(batch, name, ext, data, user))
        except Exception as exc:  # disk full, permissions... — still don't lose the rest of the batch
            reason = f"Could not store file: {exc}"
            rejected.append({"filename": name, "reason": reason})
            results.append({"filename": name, "status": "rejected", "reason": reason})

    if rejected:
        batches.update_one({"batch_id": batch_id}, {"$push": {"rejected": {"$each": rejected}}})
    return {"results": results}


def _store_resume(batch: dict, name: str, ext: str, data: bytes, user: str) -> dict:
    digest = hashlib.sha256(data).hexdigest()
    job_ids = batch.get("job_ids") or []
    now = datetime.utcnow()

    # Byte-identical file already in the database: don't re-parse (saves AI cost), just make sure
    # the original gets scored for the jobs chosen in this upload.
    existing = resumes.find_one({"sha256": digest, "status": {"$in": ["queued", "processing", "parsed"]},
                                 "duplicate_of": {"$exists": False}})
    rid = ObjectId()
    doc = {"_id": rid, "filename": name, "size": len(data), "sha256": digest, "batch_id": batch["batch_id"],
           "uploaded_by": user, "uploaded_at": now, "tags": [], "notes": []}
    if existing:
        doc.update({"status": "duplicate", "duplicate_of": str(existing["_id"]), "duplicate_confidence": "high",
                    "duplicate_reason": "identical file already uploaded", "name": existing.get("name") or name})
        resumes.insert_one(doc)
        for jid in job_ids:
            if existing["status"] == "parsed":
                if not scores.find_one({"job_id": jid, "resume_id": str(existing["_id"]), "status": "scored"}):
                    worker.enqueue("score", existing["_id"], jid, batch["batch_id"])
            else:
                resumes.update_one({"_id": existing["_id"]}, {"$addToSet": {"pending_job_ids": jid}})
        return {"filename": name, "status": "duplicate", "resume_id": str(rid),
                "reason": "Identical file already uploaded — skipped re-processing"}

    folder = os.path.join(config.upload_dir(), batch["batch_id"])
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"{rid}{ext}")
    with open(path, "wb") as fh:
        fh.write(data)
    doc.update({"status": "queued", "file_path": path, "pending_job_ids": list(job_ids)})
    resumes.insert_one(doc)
    worker.enqueue("parse", rid, None, batch["batch_id"])
    return {"filename": name, "status": "queued", "resume_id": str(rid)}


@router.get("/batches")
def recent_batches(limit: int = Query(10, ge=1, le=50), _: str = Depends(require_access)):
    out = []
    for b in batches.find().sort("created_at", -1).limit(limit):
        out.append({"batch_id": b["batch_id"], "created_at": _iso(b.get("created_at")), "created_by": b.get("created_by"),
                    "job_ids": b.get("job_ids", []), **_batch_progress(b)})
    return {"batches": out}


@router.get("/batches/{batch_id}")
def batch_status(batch_id: str, _: str = Depends(require_access)):
    b = batches.find_one({"batch_id": batch_id})
    if not b:
        raise HTTPException(404, "Upload batch not found")
    return {"batch_id": batch_id, "job_ids": b.get("job_ids", []), **_batch_progress(b, with_failures=True)}


def _batch_progress(b: dict, with_failures: bool = False) -> dict:
    bid = b["batch_id"]
    counts = {"queued": 0, "processing": 0, "parsed": 0, "failed": 0, "duplicate": 0}
    for row in resumes.aggregate([{"$match": {"batch_id": bid}}, {"$group": {"_id": "$status", "n": {"$sum": 1}}}]):
        counts[row["_id"]] = row["n"]
    rejected = b.get("rejected", [])
    received = sum(counts.values())
    sc = {"total": 0, "pending": 0, "failed": 0}
    for row in tasks.aggregate([{"$match": {"batch_id": bid, "kind": "score"}},
                                {"$group": {"_id": "$status", "n": {"$sum": 1}}}]):
        sc["total"] += row["n"]
        if row["_id"] in ("queued", "running"):
            sc["pending"] += row["n"]
        elif row["_id"] == "failed":
            sc["failed"] += row["n"]
    expected = max(b.get("expected_files", 0), received + len(rejected))
    parsing_left = counts["queued"] + counts["processing"]
    # parsed resumes whose scoring tasks aren't queued yet (the worker clears pending_job_ids right after queuing them)
    handoff = resumes.count_documents({"batch_id": bid, "status": "parsed", "pending_job_ids.0": {"$exists": True}})
    finished = (received + len(rejected) >= expected) and parsing_left == 0 and sc["pending"] == 0 and handoff == 0
    out = {"expected": expected, "received": received, "rejected_count": len(rejected),
           "counts": counts, "scoring": sc, "finished": finished}
    if with_failures:
        out["rejected"] = rejected
        out["failures"] = [{"resume_id": str(r["_id"]), "filename": r.get("filename"), "error": r.get("error")}
                           for r in resumes.find({"batch_id": bid, "status": "failed"},
                                                 {"filename": 1, "error": 1}).limit(500)]
    return out


@router.post("/batches/{batch_id}/retry")
def retry_failed(batch_id: str, _: str = Depends(require_access)):
    _require_ai()
    n = 0
    for r in resumes.find({"batch_id": batch_id, "status": "failed"}, {"_id": 1}):
        resumes.update_one({"_id": r["_id"]}, {"$set": {"status": "queued", "error": None}})
        worker.enqueue("parse", r["_id"], None, batch_id)
        n += 1
    return {"requeued": n}


# ─────────────────────────────────────────────────────────────────────────────
# Resume list / search / ranking
# ─────────────────────────────────────────────────────────────────────────────

def _resume_row(r: dict) -> dict:
    p = r.get("parsed") or {}
    latest = (p.get("experience") or [{}])[0] if p.get("experience") else {}
    return {
        "id": str(r["_id"]), "filename": r.get("filename"), "status": r.get("status"), "error": r.get("error"),
        "name": r.get("name"), "email": r.get("email"), "phone": r.get("phone"), "location": r.get("location"),
        "headline": p.get("headline"), "current_title": latest.get("title"), "current_company": latest.get("company"),
        "total_years": r.get("total_years"), "skills": (p.get("skills") or [])[:8],
        "education": EDU_LABELS.get(p.get("highest_education_level") or "none"),
        "uploaded_at": _iso(r.get("uploaded_at")), "tags": r.get("tags", []), "notes_count": len(r.get("notes", [])),
        "duplicate_of": r.get("duplicate_of"), "duplicate_reason": r.get("duplicate_reason"),
        "duplicate_confidence": r.get("duplicate_confidence"),
        # Free, job-independent ranking score computed at parse time (see parser.general_score) —
        # this is what "AI ranks every resume, top shows first" runs on when no job is picked.
        "general_score": r.get("general_score"), "general_breakdown": r.get("general_breakdown"),
    }


def _score_view(s: dict, job: Optional[dict] = None) -> dict:
    out = {"job_id": s["job_id"], "status": s.get("status"), "error": s.get("error"), "total": s.get("total"),
           "stage": s.get("stage", "new"), "flags": s.get("flags", []), "summary": s.get("summary"),
           "strengths": s.get("strengths", []), "concerns": s.get("concerns", []),
           "sub_scores": s.get("sub_scores", {}), "capped": s.get("capped", False), "masked": s.get("masked", False),
           "meets_all_must_haves": s.get("meets_all_must_haves"), "scored_at": _iso(s.get("scored_at")),
           "model": s.get("model")}
    if job is not None:
        out["stale"] = s.get("criteria_sig") != job.get("criteria_sig")
        out["job_title"] = job.get("title")
    return out


def _date(value: Optional[str], end: bool = False) -> Optional[datetime]:
    if not value:
        return None
    try:
        d = datetime.strptime(value, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(422, f"Bad date '{value}' (use YYYY-MM-DD)")
    return d + timedelta(days=1) if end else d


class Filters(BaseModel):
    q: Optional[str] = None
    location: Optional[str] = None      # city
    state: Optional[str] = None
    skills: Optional[str] = None
    skills_mode: str = "any"
    min_exp: Optional[float] = None
    max_exp: Optional[float] = None
    min_education: Optional[str] = None
    education: Optional[str] = None
    certification: Optional[str] = None
    job_id: Optional[str] = None
    min_score: Optional[float] = None
    max_score: Optional[float] = None
    stage: Optional[str] = None
    tag: Optional[str] = None
    status: Optional[str] = None
    batch_id: Optional[str] = None
    uploaded_from: Optional[str] = None
    uploaded_to: Optional[str] = None
    hide_duplicates: bool = True
    must_haves_ok: bool = False
    sort: Optional[str] = None
    order: str = "desc"
    ai_rank: bool = False


def _build_adhoc_job(f: Filters) -> dict:
    tokens = []
    if f.q:
        for m in re.findall(r'"([^"]+)"|(\S+)', f.q):
            tok = (m[0] or m[1]).strip()
            if tok:
                tokens.append(tok)

    skill_items = [s.strip() for s in (f.skills or "").split(",") if s.strip()]

    combined_skills = []
    seen = set()
    for s in skill_items:
        if s.lower() not in seen:
            seen.add(s.lower())
            combined_skills.append({"name": s, "must_have": f.skills_mode == "all"})
    for t in tokens:
        if t.lower() not in seen and len(t) > 1:
            seen.add(t.lower())
            combined_skills.append({"name": t, "must_have": False})

    cert = (f.certification or "").strip()
    edu_field = (f.education or "").strip()
    place = ", ".join(x.strip() for x in (f.location, f.state) if x and x.strip())

    title = f.q or (", ".join(skill_items) if skill_items else (cert or edu_field or place or "Instant Candidate Ranking"))
    desc = f"Search query / target: {f.q or 'None'}. Skills: {', '.join(skill_items) if skill_items else 'None'}."
    if cert:
        desc += f" Certification: {cert}."
    if edu_field:
        desc += f" Education / field of study: {edu_field}."
    if place:
        desc += f" Location: {place}."
    if f.min_exp:
        desc += f" Min experience: {f.min_exp} years."
    if f.max_exp:
        desc += f" Max experience: {f.max_exp} years."

    # Only the criteria actually searched for are scored (effective_weights ignores the rest and
    # renormalises), so a certification-only search is judged on the certification alone.
    weights = {"skills": 40, "experience": 25, "education": 15, "certifications": 25, "location": 10}

    job_spec = {
        "title": title[:120],
        "department": "Instant AI Ranking",
        "description": desc,
        "required_skills": combined_skills,
        "preferred_skills": [],
        "min_years": f.min_exp or 0.0,
        "max_years": f.max_exp,
        "education_level": f.min_education or "none",
        "education_fields": [edu_field] if edu_field else [],
        "required_certifications": [{"name": cert, "must_have": False}] if cert else [],
        "preferred_certifications": [],
        "industries": [],
        "location": place or None,
        "location_type": "onsite",
        "weights": weights,
        "mask_personal_info": False,
        "must_have_action": "flag",
        "must_have_cap": 50,
    }
    job_spec["criteria_sig"] = criteria_signature(job_spec)
    return job_spec


def _heuristic_adhoc_score(job: dict, parsed: dict) -> dict:
    weights = effective_weights(job)
    sub = {}
    flags = []

    req_skills = job.get("required_skills") or []
    cand_skills_raw = (parsed.get("skills") or [])
    cand_skills = [s.lower() for s in cand_skills_raw]
    cand_text = ((parsed.get("headline") or "") + " " + " ".join(cand_skills) + " " + (parsed.get("name") or "")).lower()

    matched = []
    missing = []
    for r in req_skills:
        name = r["name"]
        n_low = name.lower()
        found = False
        if any(n_low in cs or cs in n_low for cs in cand_skills):
            found = True
        elif n_low in cand_text:
            found = True
        elif n_low in ("react", "reactjs", "react.js") and any("react" in cs for cs in cand_skills):
            found = True
        elif n_low in ("sql", "mysql", "postgres", "postgresql") and any("sql" in cs for cs in cand_skills):
            found = True
        elif n_low in ("tally", "tally erp", "tally prime", "tally erp 9") and any("tally" in cs for cs in cand_skills):
            found = True
        elif n_low in ("gst", "goods and services tax") and any("gst" in cs for cs in cand_skills):
            found = True
        elif n_low in ("audit", "statutory audit", "internal audit", "tax audit") and any("audit" in cs for cs in cand_skills):
            found = True
        elif n_low in ("ca", "chartered accountant") and any(x in cand_text for x in ("ca", "chartered accountant")):
            found = True

        if found:
            matched.append({"name": name, "match": "full", "evidence": f"Found in profile: {name}"})
        else:
            missing.append({"name": name, "match": "none", "evidence": ""})
            if r.get("must_have"):
                flags.append({"type": "missing_skill", "severity": "critical", "message": f"Missing required skill: {name}"})

    if req_skills:
        skill_score = round(100.0 * len(matched) / len(req_skills), 1)
        sub["skills"] = {
            "score": skill_score,
            "weight": weights.get("skills", 50),
            "matched": matched,
            "partial": [],
            "missing": missing,
            "preferred_matched": [],
            "preferred_missing": []
        }

    total_y = parsed.get("total_years_experience") or 0.0
    min_y = float(job.get("min_years") or 0)
    if "experience" in weights:
        exp_score = 100.0 if min_y <= 0 else min(100.0, round(100.0 * total_y / min_y, 1))
        sub["experience"] = {
            "score": exp_score,
            "weight": weights["experience"],
            "relevant_years": total_y,
            "total_years": total_y,
            "required_years": min_y,
            "note": f"{total_y:g} years experience vs {min_y:g} requested"
        }

    need_rank = EDU_RANKS.get(job.get("education_level") or "none", 0)
    have_rank = EDU_RANKS.get(parsed.get("highest_education_level") or "none", 0)
    if "education" in weights:
        edu_score = 100.0 if have_rank >= need_rank else max(30.0, round(100.0 * have_rank / max(need_rank, 1), 1))
        sub["education"] = {
            "score": edu_score,
            "weight": weights["education"],
            "highest": parsed.get("highest_education_level") or "none",
            "required": job.get("education_level") or "none"
        }

    if "certifications" in weights:
        have_certs = [str(c).lower() for c in (parsed.get("certifications") or [])]
        want = [c["name"] if isinstance(c, dict) else str(c) for c in (job.get("required_certifications") or [])]
        got = [{"name": w, "match": "full", "evidence": "Listed in the CV's certifications"}
               for w in want if any(w.lower() in hc or hc in w.lower() for hc in have_certs)]
        lost = [{"name": w, "match": "none", "evidence": ""} for w in want if w not in [g["name"] for g in got]]
        sub["certifications"] = {"score": round(100.0 * len(got) / len(want), 1) if want else 0.0,
                                 "weight": weights["certifications"], "matched": got, "missing": lost}

    if "location" in weights:
        loc_have = (parsed.get("location") or "").lower()
        parts = [x.strip().lower() for x in (job.get("location") or "").split(",") if x.strip()]
        hit = [x for x in parts if x in loc_have]
        sub["location"] = {"score": round(100.0 * len(hit) / len(parts), 1) if parts else 0.0,
                           "weight": weights["location"], "note": f"CV location: {parsed.get('location') or 'not stated'}"}

    total = sum(sub[c]["score"] * sub[c]["weight"] for c in sub) / sum(sub[c]["weight"] for c in sub) if sub else 75.0

    strengths = []
    if matched:
        strengths.append(f"Matched skills: {', '.join(m['name'] for m in matched[:3])}")
    if total_y > 0:
        strengths.append(f"{total_y:g} years professional experience")
    if "education" in sub and have_rank >= 3:
        strengths.append(f"Education: {EDU_LABELS.get(parsed.get('highest_education_level'), parsed.get('highest_education_level'))}")
    if "certifications" in sub and sub["certifications"]["matched"]:
        strengths.append(f"Certification: {', '.join(m['name'] for m in sub['certifications']['matched'])}")

    concerns = []
    if missing:
        concerns.append(f"Not explicitly matched: {', '.join(m['name'] for m in missing[:3])}")
    if "experience" in sub and min_y > 0 and total_y < min_y:
        concerns.append(f"Experience ({total_y:g}y) is below preferred minimum ({min_y:g}y)")

    return {
        "total": round(total, 1),
        "sub_scores": sub,
        "flags": flags,
        "capped": False,
        "meets_all_must_haves": not any(f["severity"] == "critical" for f in flags),
        "summary": f"Candidate has {total_y:g} years experience with skills in {', '.join(cand_skills_raw[:4])}. Meets {len(matched)} of {len(req_skills)} requested target skills/keywords.",
        "strengths": strengths[:4],
        "concerns": concerns[:4],
        "model": "heuristic-fallback"
    }


def _score_one_adhoc(job_spec: dict, doc: dict) -> dict:
    sig = job_spec.get("criteria_sig") or "sig"
    cache_key = f"{sig}_{doc['_id']}"
    cached = adhoc_scores.find_one({"_id": cache_key})
    if cached and "score" in cached:
        return cached["score"]

    parsed = doc.get("parsed") or {}
    try:
        if llm.is_configured():
            result = score_resume(job_spec, parsed)
        else:
            result = _heuristic_adhoc_score(job_spec, parsed)
    except Exception as exc:
        log.warning("AI scoring failed for resume %s: %s; using heuristic fallback", doc.get("_id"), exc)
        result = _heuristic_adhoc_score(job_spec, parsed)

    result["job_id"] = f"adhoc_{sig}"
    result["job_title"] = job_spec["title"]
    adhoc_scores.update_one(
        {"_id": cache_key},
        {"$set": {"sig": sig, "resume_id": str(doc["_id"]), "score": result, "created_at": datetime.utcnow()}},
        upsert=True
    )
    return result


def _score_adhoc_resumes(job_spec: dict, docs: List[dict]) -> Dict[str, dict]:
    results = {}
    to_fetch = []
    sig = job_spec.get("criteria_sig") or "sig"

    for d in docs:
        rid = str(d["_id"])
        cache_key = f"{sig}_{d['_id']}"
        cached = adhoc_scores.find_one({"_id": cache_key})
        if cached and "score" in cached:
            results[rid] = cached["score"]
        else:
            to_fetch.append(d)

    if to_fetch:
        max_workers = min(len(to_fetch), 8)
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_doc = {executor.submit(_score_one_adhoc, job_spec, d): str(d["_id"]) for d in to_fetch}
            for future in future_to_doc:
                rid = future_to_doc[future]
                try:
                    results[rid] = future.result()
                except Exception as exc:
                    log.error("Scoring thread error for %s: %s", rid, exc)
                    results[rid] = {"total": 50.0, "status": "error", "error": str(exc), "sub_scores": {}, "flags": []}

    return results


def _pick_adhoc_candidates(job_spec: dict, docs: List[dict]) -> tuple[List[dict], List[dict]]:
    """(to_ai, to_heuristic) — hard cap on how many documents ever reach the
    real LLM scorer for one ad-hoc "AI Rank" search request (see
    config.adhoc_ai_rank_max()). A search with no job selected can match an
    arbitrarily large, unfiltered slice of the database — without this cap,
    one request could fan out a real OpenAI call per match. Beyond the cap,
    candidates still get a real (free, deterministic) heuristic score, so the
    list stays complete and sorted — just less finely judged past the top-K.
    Never trust the client-supplied ai_rank flag to self-limit this: the cap
    is enforced here regardless of how the request asked for AI ranking."""
    cap = config.adhoc_ai_rank_max()
    if len(docs) <= cap:
        return docs, []
    ids = [d["_id"] for d in docs]
    emb_map = {d["_id"]: d.get("embedding") for d in resumes.find({"_id": {"$in": ids}}, {"embedding": 1})}
    job_vec = embed_job(job_spec)
    if not job_vec:
        # No similarity signal for this query at all -> can't pick a meaningful top-K,
        # so nobody takes the expensive path (safe default: zero extra AI cost).
        return [], docs
    ranked = sorted(docs, key=lambda d: cosine_similarity(emb_map.get(d["_id"]), job_vec), reverse=True)
    return ranked[:cap], ranked[cap:]


def _assign_adhoc_ranks(rows: List[dict]) -> None:
    totals = sorted([
        r["score"]["total"] for r in rows
        if r.get("score") and r["score"].get("total") is not None and not r.get("duplicate_of")
    ], reverse=True)
    for row in rows:
        total = (row.get("score") or {}).get("total")
        row["rank"] = None if row.get("duplicate_of") or total is None else 1 + sum(1 for t in totals if t > total)


def _word_rx(term: str) -> str:
    """Whole-word match for a search term. A plain substring match made short keywords useless:
    "CA" matched "Associate", "Chartered", "Location"... so every CV matched and the list never
    changed. Boundaries are lookarounds (not \b) so terms like "C++" or ".NET" still work."""
    return r"(?<![A-Za-z0-9])" + re.escape(term.strip()) + r"(?![A-Za-z0-9])"


def _search(f: Filters) -> List[dict]:
    """Every match (unpaged), each row already carrying its score when job_id is set or ai_rank is requested."""
    flt: dict = {}
    if f.status == "all":
        pass
    elif f.status:
        flt["status"] = f.status
    else:
        # identical-file re-uploads are stored as status "duplicate" (never parsed); show them only when asked
        flt["status"] = "parsed" if f.hide_duplicates else {"$in": ["parsed", "duplicate"]}
    and_: list = []

    # Only an explicit ai_rank=true opts into AI cost — sort=score alone must never
    # silently trigger it (that used to be an unbounded-cost default: every plain
    # search/page-load could fire one real LLM call per match). The cap in
    # _pick_adhoc_candidates is a second, independent safety net on top of this.
    should_ai_rank = not f.job_id and f.ai_rank

    if f.q:
        tokens = [tok[0] or tok[1] for tok in re.findall(r'"([^"]+)"|(\S+)', f.q)]
        if should_ai_rank and f.ai_rank and len(tokens) > 1:
            or_terms = [{"$or": [{"text": {"$regex": _word_rx(t), "$options": "i"}},
                                 {"name": {"$regex": _word_rx(t), "$options": "i"}},
                                 {"filename": {"$regex": _word_rx(t), "$options": "i"}}]}
                        for t in tokens if t]
            if or_terms:
                and_.append({"$or": or_terms})
        else:
            for tok in tokens:
                term = _word_rx(tok)
                rx = {"$regex": term, "$options": "i"}
                and_.append({"$or": [{"text": rx}, {"name": rx}, {"filename": rx}]})

    if f.skills and not (should_ai_rank and f.ai_rank):
        keys = [k for k in (skill_key(s) for s in f.skills.split(",")) if k]
        if keys:
            flt["skills_key"] = {"$all": keys} if f.skills_mode == "all" else {"$in": keys}
    elif f.skills and should_ai_rank and f.ai_rank:
        keys = [k for k in (skill_key(s) for s in f.skills.split(",")) if k]
        if keys:
            flt["skills_key"] = {"$all": keys} if f.skills_mode == "all" else {"$in": keys}

    if f.min_exp is not None or f.max_exp is not None:
        rng = {}
        if f.min_exp is not None:
            rng["$gte"] = f.min_exp
        if f.max_exp is not None:
            rng["$lte"] = f.max_exp
        flt["total_years"] = rng
    if f.min_education:
        if f.min_education not in EDU_RANKS:
            raise HTTPException(422, "Unknown education level")
        flt["education_rank"] = {"$gte": EDU_RANKS[f.min_education]}
    if f.location and f.location.strip():      # city
        flt["location"] = {"$regex": _word_rx(f.location.strip()), "$options": "i"}
    if f.state and f.state.strip():            # state — ANDed with the city, both are in the one location string
        and_.append({"location": {"$regex": _word_rx(f.state.strip()), "$options": "i"}})
    if f.education and f.education.strip():
        erx = {"$regex": _word_rx(f.education.strip()), "$options": "i"}
        and_.append({"$or": [{"parsed.education.degree": erx}, {"parsed.education.field": erx}]})
    if f.certification and f.certification.strip():
        and_.append({"parsed.certifications": {"$regex": _word_rx(f.certification.strip()), "$options": "i"}})
    if f.tag:
        flt["tags"] = f.tag.strip().lower()
    if f.batch_id:
        flt["batch_id"] = f.batch_id
    d_from, d_to = _date(f.uploaded_from), _date(f.uploaded_to, end=True)
    if d_from or d_to:
        flt["uploaded_at"] = {**({"$gte": d_from} if d_from else {}), **({"$lt": d_to} if d_to else {})}
    if f.hide_duplicates:
        and_.append(NOT_HARD_DUP)
    if and_:
        flt["$and"] = and_

    score_by_resume: dict = {}
    job = None
    adhoc_job = None
    if f.job_id:
        job = _get_job(f.job_id)
        # screened_out candidates are included by default (visible, unranked, clearly labelled —
        # see prefilter.py) but excluded the moment a score/must-haves filter is applied, since a
        # screened-out record has no score to filter by.
        sflt: dict = {"job_id": f.job_id, "status": {"$in": ["scored", "screened_out"]}}
        if f.min_score is not None or f.max_score is not None:
            sflt["status"] = "scored"
            sflt["total"] = {**({"$gte": f.min_score} if f.min_score is not None else {}),
                             **({"$lte": f.max_score} if f.max_score is not None else {})}
        if f.stage:
            sflt["stage"] = f.stage
        if f.must_haves_ok:
            sflt["meets_all_must_haves"] = True
        score_by_resume = {s["resume_id"]: s for s in scores.find(sflt)}
        flt["_id"] = {"$in": [ObjectId(i) for i in score_by_resume]}

    docs = list(resumes.find(flt, LIST_EXCLUDE))

    # If no job is set and AI ranking is requested, compute AI scores across the matched documents —
    # capped (see _pick_adhoc_candidates) so this can never mean one real LLM call per match.
    if not job and should_ai_rank and docs:
        adhoc_job = _build_adhoc_job(f)
        to_ai, to_heuristic = _pick_adhoc_candidates(adhoc_job, docs)
        score_by_resume = _score_adhoc_resumes(adhoc_job, to_ai)
        for d in to_heuristic:
            result = _heuristic_adhoc_score(adhoc_job, d.get("parsed") or {})
            result["job_id"] = f"adhoc_{adhoc_job.get('criteria_sig') or 'sig'}"
            result["job_title"] = adhoc_job["title"]
            score_by_resume[str(d["_id"])] = result

    # identical-file re-uploads are never parsed themselves: show the original's details
    canon_ids = {ObjectId(d["duplicate_of"]) for d in docs if d.get("status") == "duplicate" and d.get("duplicate_of")}
    canon = {c["_id"]: _resume_row(c) for c in resumes.find({"_id": {"$in": list(canon_ids)}}, LIST_EXCLUDE)}
    own = {"id", "filename", "status", "error", "uploaded_at", "tags", "notes_count", "duplicate_of",
           "duplicate_reason", "duplicate_confidence"}
    rows = []
    for r in docs:
        row = _resume_row(r)
        src = canon.get(ObjectId(r["duplicate_of"])) if r.get("status") == "duplicate" and r.get("duplicate_of") else None
        if src:
            row.update({k: v for k, v in src.items() if k not in own})
        if job is not None:
            row["score"] = _score_view(score_by_resume[row["id"]], job)
        elif adhoc_job is not None and row["id"] in score_by_resume:
            row["score"] = _score_view(score_by_resume[row["id"]], adhoc_job)
        rows.append(row)

    def _score_of(row: dict):
        """The real job/AI score when there is one, else the free general_score —
        so "top-ranked first" always means something, with or without a job picked."""
        sc = row.get("score")
        return sc.get("total") if sc is not None else row.get("general_score")

    if job is None:
        # Job-scoped min/max-score filtering already happened in the sflt query above
        # (cheaper: filters before fetching). Here it's general_score/adhoc-score, computed
        # in Python, so it's applied after rows are built.
        if f.min_score is not None or f.max_score is not None:
            def _score_ok(row):
                tot = _score_of(row)
                if tot is None: return False
                if f.min_score is not None and tot < f.min_score: return False
                if f.max_score is not None and tot > f.max_score: return False
                return True
            rows = [r for r in rows if _score_ok(r)]
        if f.must_haves_ok:
            rows = [r for r in rows if (r.get("score") or {}).get("meets_all_must_haves")]

    # Default sort is always "top-ranked first" — the free general_score when no job/AI-rank
    # score is present, so a plain search with zero clicks still shows the best resumes first.
    key = f.sort or "score"
    reverse = f.order != "asc"
    sorters = {
        "score": lambda x: (_score_of(x) if _score_of(x) is not None else -1, (x.get("name") or "").lower()),
        "name": lambda x: (x.get("name") or "").lower(),
        "experience": lambda x: x.get("total_years") if x.get("total_years") is not None else -1,
        "uploaded": lambda x: x.get("uploaded_at") or "",
    }
    rows.sort(key=sorters.get(key, sorters["score"]), reverse=reverse)
    if job is not None:
        _assign_ranks(rows, f.job_id)
    elif adhoc_job is not None:
        _assign_adhoc_ranks(rows)
    else:
        _assign_general_ranks(rows)
    return rows


def _assign_general_ranks(rows: List[dict]) -> None:
    """Rank by the free general_score when no job and no AI-rank search is active —
    this is the default "AI ranks every resume" view (see parser.general_score)."""
    totals = sorted((r["general_score"] for r in rows if r.get("general_score") is not None and not r.get("duplicate_of")),
                    reverse=True)
    for row in rows:
        total = row.get("general_score")
        row["rank"] = None if row.get("duplicate_of") or total is None else 1 + sum(1 for t in totals if t > total)


def _assign_ranks(rows: List[dict], job_id: str) -> None:
    """Rank = position among ALL scored, non-duplicate candidates of the job (so it
    stays meaningful when filters narrow the visible list). Ties share a rank."""
    all_scores = list(scores.find({"job_id": job_id, "status": "scored"}, {"resume_id": 1, "total": 1}))
    dup_ids = {str(d["_id"]) for d in resumes.find(
        {"_id": {"$in": [ObjectId(s["resume_id"]) for s in all_scores]}, "duplicate_of": {"$exists": True}, "duplicate_confidence": "high"}, {"_id": 1})}
    totals = sorted((s["total"] for s in all_scores if s["resume_id"] not in dup_ids), reverse=True)
    for row in rows:
        total = (row.get("score") or {}).get("total")
        row["rank"] = None if row["id"] in dup_ids or total is None else 1 + sum(1 for t in totals if t > total)


@router.get("/locations")
def list_locations(_: str = Depends(require_access)):
    """Distinct candidate locations found in uploaded CVs (case-insensitive), with CV counts,
    for the Location dropdown. Duplicates and unparsed files are left out."""
    seen: dict = {}
    for d in resumes.find({"status": "parsed", **NOT_HARD_DUP, "location": {"$nin": [None, ""]}}, {"location": 1}):
        loc = " ".join(str(d.get("location") or "").split())
        if not loc:
            continue
        entry = seen.setdefault(loc.lower(), {"location": loc, "count": 0})
        entry["count"] += 1
    return {"locations": sorted(seen.values(), key=lambda e: e["location"].lower())}


@router.get("/education-options")
def list_education_options(_: str = Depends(require_access)):
    """Distinct degrees / fields of study found in uploaded CVs, with CV counts, for the Education dropdown."""
    seen: dict = {}
    for d in resumes.find({"status": "parsed", **NOT_HARD_DUP}, {"parsed.education": 1}):
        names = set()
        for e in (d.get("parsed") or {}).get("education") or []:
            for key in ("degree", "field"):
                v = " ".join(str(e.get(key) or "").split())
                if v:
                    names.add(v)
        for v in names:
            entry = seen.setdefault(v.lower(), {"name": v, "count": 0})
            entry["count"] += 1
    return {"education": sorted(seen.values(), key=lambda e: e["name"].lower())}


@router.get("/certification-options")
def list_certification_options(_: str = Depends(require_access)):
    """Distinct certifications found in uploaded CVs, with CV counts, for the Certification dropdown."""
    seen: dict = {}
    for d in resumes.find({"status": "parsed", **NOT_HARD_DUP}, {"parsed.certifications": 1}):
        names = {" ".join(str(c or "").split()) for c in (d.get("parsed") or {}).get("certifications") or []}
        for v in names:
            if v:
                entry = seen.setdefault(v.lower(), {"name": v, "count": 0})
                entry["count"] += 1
    return {"certifications": sorted(seen.values(), key=lambda e: e["name"].lower())}


@router.get("/resumes")
def list_resumes(
    q: Optional[str] = None, location: Optional[str] = None, state: Optional[str] = None, skills: Optional[str] = None, skills_mode: str = "any",
    education: Optional[str] = None, certification: Optional[str] = None,
    min_exp: Optional[float] = None, max_exp: Optional[float] = None, min_education: Optional[str] = None,
    job_id: Optional[str] = None, min_score: Optional[float] = None, max_score: Optional[float] = None,
    stage: Optional[str] = None, tag: Optional[str] = None, status: Optional[str] = None,
    batch_id: Optional[str] = None, uploaded_from: Optional[str] = None, uploaded_to: Optional[str] = None,
    hide_duplicates: bool = True, must_haves_ok: bool = False,
    sort: Optional[str] = None, order: str = "desc", ai_rank: bool = False,
    page: int = Query(1, ge=1), page_size: int = Query(25, ge=1, le=100),
    _: str = Depends(require_access),
):
    f = Filters(q=q, location=location, state=state, skills=skills, skills_mode=skills_mode, min_exp=min_exp, max_exp=max_exp,
                education=education, certification=certification, min_education=min_education, job_id=job_id, min_score=min_score, max_score=max_score, stage=stage,
                tag=tag, status=status, batch_id=batch_id, uploaded_from=uploaded_from, uploaded_to=uploaded_to,
                hide_duplicates=hide_duplicates, must_haves_ok=must_haves_ok, sort=sort, order=order, ai_rank=ai_rank)
    rows = _search(f)
    start = (page - 1) * page_size
    out = {"total": len(rows), "page": page, "page_size": page_size, "resumes": rows[start:start + page_size]}
    if job_id:
        parsed_ids = {str(d["_id"]) for d in resumes.find({"status": "parsed", **NOT_HARD_DUP}, {"_id": 1})}
        handled = {s["resume_id"] for s in scores.find({"job_id": job_id}, {"resume_id": 1})}
        out["unscored"] = len(parsed_ids - handled)
        out["errors"] = scores.count_documents({"job_id": job_id, "status": "error"})
    return out


@router.get("/resumes/{resume_id}")
def resume_detail(resume_id: str, _: str = Depends(require_access)):
    r = _get_resume(resume_id)
    row = _resume_row(r)
    row.update({"text": r.get("text"), "parsed": r.get("parsed"), "size": r.get("size"),
                "notes": [{**n, "at": _iso(n.get("at"))} for n in r.get("notes", [])]})
    job_map = {str(j["_id"]): j for j in jobs.find()}
    row["scores"] = [_score_view(s, job_map[s["job_id"]]) for s in scores.find({"resume_id": resume_id})
                     if s["job_id"] in job_map]
    for a in adhoc_scores.find({"resume_id": resume_id}).sort("created_at", -1).limit(5):
        sc = a.get("score") or {}
        row["scores"].append({
            "job_id": sc.get("job_id", str(a["_id"])),
            "job_title": sc.get("job_title", "AI Rank evaluation"),
            **sc
        })
    row["scores"].sort(key=lambda s: -(s.get("total") or -1))
    if r.get("duplicate_of"):
        canon = resumes.find_one({"_id": _oid(r["duplicate_of"])}, {"name": 1, "filename": 1})
        row["duplicate_of_name"] = (canon or {}).get("name") or (canon or {}).get("filename")
    else:
        row["duplicates"] = [{"id": str(d["_id"]), "filename": d.get("filename"), "reason": d.get("duplicate_reason"),
                              "confidence": d.get("duplicate_confidence")}
                             for d in resumes.find({"duplicate_of": resume_id}, {"filename": 1, "duplicate_reason": 1,
                                                                                "duplicate_confidence": 1})]
    return row


@router.get("/resumes/{resume_id}/file")
def download_original(resume_id: str, inline: bool = False, _: str = Depends(require_access)):
    """The CV exactly as uploaded. inline=true lets the browser display it (PDF) instead of saving it."""
    r = _get_resume(resume_id)
    path = r.get("file_path")
    if not path or not os.path.exists(path):
        raise HTTPException(404, "Original file is not available")
    name = r.get("filename") or "resume"
    return FileResponse(path, filename=name, media_type=mimetypes.guess_type(name)[0] or "application/octet-stream",
                        content_disposition_type="inline" if inline else "attachment")


def _delete_resume_doc(r: dict) -> None:
    """Remove one resume everywhere it lives: stored file, scores (job + cached AI Rank), queued
    tasks, and any duplicate links pointing at it."""
    rid = str(r["_id"])
    if r.get("file_path"):
        try:
            os.remove(r["file_path"])
        except OSError:
            pass
    scores.delete_many({"resume_id": rid})
    adhoc_scores.delete_many({"resume_id": rid})
    tasks.delete_many({"resume_id": rid, "status": "queued"})
    resumes.update_many({"duplicate_of": rid}, {"$unset": {"duplicate_of": "", "duplicate_reason": "",
                                                           "duplicate_confidence": ""}})
    resumes.delete_one({"_id": r["_id"]})


@router.delete("/resumes/{resume_id}")
def delete_resume(resume_id: str, user: str = Depends(require_access)):
    r = _get_resume(resume_id)
    _delete_resume_doc(r)
    log.info("resume %s (%s) deleted by %s", resume_id, r.get("filename"), user)
    return {"success": True}


class BulkDeleteIn(BaseModel):
    ids: List[str] = Field(min_length=1, max_length=200)


@router.post("/resumes/bulk-delete")
def bulk_delete_resumes(body: BulkDeleteIn, user: str = Depends(require_access)):
    """Delete several CVs at once (the selected rows). Unknown ids are ignored."""
    oids = [_oid(i, "Resume") for i in body.ids]
    deleted = 0
    for r in resumes.find({"_id": {"$in": oids}}):
        _delete_resume_doc(r)
        deleted += 1
    log.info("%d resume(s) bulk-deleted by %s", deleted, user)
    return {"deleted": deleted}


@router.post("/resumes/{resume_id}/reparse")
def reparse(resume_id: str, _: str = Depends(require_access)):
    r = _get_resume(resume_id)
    _require_ai()
    if not r.get("file_path"):
        raise HTTPException(409, "This entry has no stored file to re-read")
    jids = [s["job_id"] for s in scores.find({"resume_id": resume_id}, {"job_id": 1})]
    resumes.update_one({"_id": r["_id"]}, {"$set": {"status": "queued", "error": None, "pending_job_ids": jids},
                                           "$unset": {"text": ""}})
    worker.enqueue("parse", r["_id"])
    return {"success": True}


class TagsIn(BaseModel):
    tags: List[str]


@router.put("/resumes/{resume_id}/tags")
def set_tags(resume_id: str, body: TagsIn, _: str = Depends(require_access)):
    r = _get_resume(resume_id, {"_id": 1})
    tags = []
    for t in body.tags:
        t = re.sub(r"\s+", " ", t).strip().lower()[:30]
        if t and t not in tags:
            tags.append(t)
    resumes.update_one({"_id": r["_id"]}, {"$set": {"tags": tags[:20]}})
    return {"tags": tags[:20]}


class NoteIn(BaseModel):
    text: str = Field(min_length=1, max_length=2000)


@router.post("/resumes/{resume_id}/notes")
def add_note(resume_id: str, body: NoteIn, user: str = Depends(require_access)):
    r = _get_resume(resume_id, {"_id": 1})
    note = {"id": secrets.token_hex(6), "text": body.text.strip(), "author": user, "at": datetime.utcnow()}
    resumes.update_one({"_id": r["_id"]}, {"$push": {"notes": note}})
    return {"note": {**note, "at": _iso(note["at"])}}


@router.delete("/resumes/{resume_id}/notes/{note_id}")
def delete_note(resume_id: str, note_id: str, _: str = Depends(require_access)):
    r = _get_resume(resume_id, {"_id": 1})
    resumes.update_one({"_id": r["_id"]}, {"$pull": {"notes": {"id": note_id}}})
    return {"success": True}


@router.get("/tags")
def all_tags(_: str = Depends(require_access)):
    return {"tags": sorted(t for t in resumes.distinct("tags") if t)}


@router.get("/overview")
def overview(_: str = Depends(require_access)):
    counts = {row["_id"]: row["n"] for row in resumes.aggregate([{"$group": {"_id": "$status", "n": {"$sum": 1}}}])}
    return {"resumes": counts, "total_resumes": sum(counts.values()), "jobs": jobs.count_documents({}),
            "queue": {"queued": tasks.count_documents({"status": "queued"}),
                      "running": tasks.count_documents({"status": "running"})},
            "ai_configured": llm.is_configured(), "education_levels": EDU_LABELS,
            "criteria_labels": CRITERION_LABELS, "default_weights": DEFAULT_WEIGHTS}


# ─────────────────────────────────────────────────────────────────────────────
# Export
# ─────────────────────────────────────────────────────────────────────────────

def _csv_safe(v):
    """Names/skills come from untrusted documents: neutralise spreadsheet formulas."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


@router.get("/jobs/{job_id}/export")
def export_job(
    job_id: str, format: str = Query("xlsx", pattern="^(csv|xlsx)$"),
    stage: Optional[str] = None, min_score: Optional[float] = None, top: Optional[int] = Query(None, ge=1),
    ids: Optional[str] = None, include_duplicates: bool = False, must_haves_ok: bool = False,
    _: str = Depends(require_access),
):
    """Ranked candidates for a job. Filters: stage (e.g. shortlisted), min_score, top N, explicit ids."""
    job = _get_job(job_id)
    rows = _search(Filters(job_id=job_id, stage=stage, min_score=min_score, hide_duplicates=not include_duplicates,
                           must_haves_ok=must_haves_ok, sort="score"))
    if ids:
        wanted = {i.strip() for i in ids.split(",") if i.strip()}
        rows = [r for r in rows if r["id"] in wanted]
    if top:
        rows = rows[:top]
    details = {str(d["_id"]): d for d in resumes.find({"_id": {"$in": [ObjectId(r["id"]) for r in rows]}},
                                                      {"parsed.certifications": 1, "parsed.highest_education_level": 1})}
    crit = [c for c in CRITERIA if c in active_criteria(job)]
    header = (["Rank", "Name", "Email", "Phone", "Location", "Total score"]
              + [f"{CRITERION_LABELS[c]} score" for c in crit]
              + ["Total years", "Highest education", "Top skills", "Certifications", "Missing must-haves",
                 "Stage", "Tags", "Possible duplicate", "AI summary", "Uploaded", "File"])
    table = []
    for r in rows:
        s = r["score"]
        table.append([
            r.get("rank"), r["name"], r["email"], r["phone"], r["location"], s["total"],
            *[(s["sub_scores"].get(c) or {}).get("score") for c in crit],
            r["total_years"], r["education"], ", ".join(r["skills"]),
            ", ".join((details.get(r["id"], {}).get("parsed") or {}).get("certifications", [])),
            "; ".join(f["message"] for f in s["flags"] if f.get("severity") == "critical"),
            s["stage"], ", ".join(r["tags"]), r["duplicate_reason"] or "", s["summary"] or "",
            (r["uploaded_at"] or "")[:10], r["filename"],
        ])
    table = [[_csv_safe(c) for c in row] for row in table]
    fname = re.sub(r"[^A-Za-z0-9_-]+", "_", job["title"]).strip("_") or "job"
    stamp = datetime.utcnow().strftime("%Y%m%d")

    if format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(header)
        w.writerows(table)
        return Response(("﻿" + buf.getvalue()).encode("utf-8"), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="{fname}_ranking_{stamp}.csv"'})

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Ranking"
    ws.append(header)
    for row in table:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B2C6F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "C2"
    widths = [6, 26, 30, 16, 20, 11] + [12] * len(crit) + [9, 18, 40, 30, 40, 13, 20, 24, 60, 12, 30]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{fname}_ranking_{stamp}.xlsx"'})


@router.get("/export")
def export_resumes(
    q: Optional[str] = None, location: Optional[str] = None, state: Optional[str] = None, skills: Optional[str] = None, skills_mode: str = "any",
    education: Optional[str] = None, certification: Optional[str] = None,
    min_exp: Optional[float] = None, max_exp: Optional[float] = None, min_education: Optional[str] = None,
    job_id: Optional[str] = None, min_score: Optional[float] = None, max_score: Optional[float] = None,
    stage: Optional[str] = None, tag: Optional[str] = None, format: str = Query("xlsx", pattern="^(csv|xlsx)$"),
    top: Optional[int] = Query(None, ge=1), ids: Optional[str] = None, ai_rank: bool = False,
    _: str = Depends(require_access),
):
    """Export search results (no job selected) or a job's ranking (job_id set)."""
    if job_id:
        return export_job(job_id=job_id, format=format, stage=stage, min_score=min_score, top=top, ids=ids, _=_)

    f = Filters(q=q, location=location, state=state, skills=skills, skills_mode=skills_mode, min_exp=min_exp, max_exp=max_exp,
                education=education, certification=certification, min_education=min_education, min_score=min_score, max_score=max_score, stage=stage,
                tag=tag, sort="score", ai_rank=ai_rank)
    rows = _search(f)
    if ids:
        wanted = {i.strip() for i in ids.split(",") if i.strip()}
        rows = [r for r in rows if r["id"] in wanted]
    if top:
        rows = rows[:top]

    header = ["Rank", "Name", "Email", "Phone", "Location", "AI Match Score", "Total Exp (Years)",
              "Highest Education", "Top Skills", "AI Summary", "Uploaded", "File"]
    table = []
    for r in rows:
        s = r.get("score") or {}
        table.append([
            r.get("rank"), r.get("name"), r.get("email"), r.get("phone"), r.get("location"),
            s.get("total", "—"), r.get("total_years"), r.get("education"), ", ".join(r.get("skills") or []),
            s.get("summary", ""), (r.get("uploaded_at") or "")[:10], r.get("filename")
        ])
    table = [[_csv_safe(c) for c in row] for row in table]
    stamp = datetime.utcnow().strftime("%Y%m%d")

    if format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(header)
        w.writerows(table)
        return Response(("﻿" + buf.getvalue()).encode("utf-8"), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="candidate_ranking_{stamp}.csv"'})

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "AI Ranking"
    ws.append(header)
    for row in table:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B2C6F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "C2"
    widths = [8, 26, 30, 16, 20, 16, 16, 18, 40, 50, 14, 30]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="candidate_ranking_{stamp}.xlsx"'})


