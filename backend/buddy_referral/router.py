# backend/buddy_referral/router.py
"""
Buddy Referral module.

  1. Existing employee referral - an employee records which colleague (picked
     via a JHS-code / name search) referred someone, for a position.
  2. Referral - an employee submits a role plus the candidate's resume
     (.pdf / .docx only), stored in MongoDB GridFS.

Admin endpoints (control panel "Buddy Referral" card) show the submitting
employee's code, name, date of joining, manager and partner (looked up live
from Employee_details) and export the history to Excel.
"""

import re
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Optional

import gridfs
from bson import ObjectId
from bson.errors import InvalidId
from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from gridfs.errors import NoFile
from pydantic import BaseModel

from backend.access_manager.router import is_referral_admin, require_referral_admin
from backend.auth import get_current_user
from backend.buddy_referral.notify import send_referral_email
from backend.database import db, employee_details_collection

router = APIRouter(prefix="/buddy-referral", tags=["Buddy Referral"])

existing_referrals = db["buddy_existing_referrals"]
resume_referrals = db["buddy_referrals"]
_fs = gridfs.GridFS(db, collection="buddy_referral_resumes")

MAX_RESUME_BYTES = 10 * 1024 * 1024
_ALLOWED = {
    ".pdf": ("application/pdf", b"%PDF"),
    ".docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"PK\x03\x04"),
}
_IST = timedelta(hours=5, minutes=30)


# ── helpers ───────────────────────────────────────────────────────────────────
def _name(emp: dict, fallback: str = "") -> str:
    return (emp or {}).get("Emp Name") or (emp or {}).get("EmployeeName") or (emp or {}).get("Name") or fallback


def _emp(empid: str) -> Optional[dict]:
    return employee_details_collection.find_one({"EmpID": (empid or "").strip().upper()})


def _profile(emp: Optional[dict], empid: str) -> dict:
    """The employee facts the screens need (DOJ, manager, partner)."""
    emp = emp or {}
    return {
        "empid": empid, "name": _name(emp, empid),
        "date_of_joining": emp.get("Date of Joining") or "",
        "manager": emp.get("ReportingEmpName") or "",
        "partner": emp.get("Partner") or "",
        "designation": emp.get("Designation Name") or emp.get("Designation") or "",
        "branch": emp.get("Branch Name") or "",
    }


def _profiles_for(empids) -> dict:
    ids = sorted({(e or "").strip().upper() for e in empids if e})
    found = {(d.get("EmpID") or "").strip().upper(): d
             for d in employee_details_collection.find({"EmpID": {"$in": ids}})} if ids else {}
    return {e: _profile(found.get(e), e) for e in ids}


def _iso(dt):
    if isinstance(dt, datetime):
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.isoformat()
    return dt


def _ist_text(dt) -> str:
    if not isinstance(dt, datetime):
        return ""
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return (dt + _IST).strftime("%d/%m/%Y %H:%M")


# ── Signed-in employee + employee search (type-ahead) ─────────────────────────
@router.get("/me")
def me(current_user: str = Depends(get_current_user)):
    empid = current_user.strip().upper()
    return _profile(_emp(empid), empid)


@router.get("/employees/search")
def search_employees(q: str = "", current_user: str = Depends(get_current_user)):
    q = (q or "").strip()
    if len(q) < 1:
        return {"results": []}
    rx = {"$regex": re.escape(q), "$options": "i"}
    cur = employee_details_collection.find(
        {"$or": [{"EmpID": rx}, {"Emp Name": rx}, {"EmployeeName": rx}, {"Name": rx}]}
    ).limit(40)
    out = []
    for e in cur:
        empid = (e.get("EmpID") or "").strip().upper()
        if not empid:
            continue
        out.append({
            "empid": empid, "name": _name(e, empid),
            "designation": e.get("Designation Name") or e.get("Designation") or "",
            "department": e.get("Branch Name") or e.get("Department") or "",
        })
    ql = q.lower()
    out.sort(key=lambda r: (not (r["empid"].lower().startswith(ql) or r["name"].lower().startswith(ql)), r["name"].lower()))
    return {"results": out[:10]}


# ── 1. Existing employee referral ─────────────────────────────────────────────
class ExistingReferralBody(BaseModel):
    referral_from: str          # JHS code of the referring employee (selected from search)
    position: str


@router.post("/existing")
def submit_existing(body: ExistingReferralBody, background: BackgroundTasks,
                    current_user: str = Depends(get_current_user)):
    position = body.position.strip()
    if not position:
        raise HTTPException(400, "Position is required")
    ref = _emp(body.referral_from)
    if not ref:
        raise HTTPException(400, "Please select the referring employee from the suggestions")
    me_ = _profile(_emp(current_user), current_user.strip().upper())
    ref_id = (ref.get("EmpID") or "").strip().upper()
    existing_referrals.insert_one({
        "submitted_by": me_["empid"], "submitted_by_name": me_["name"],
        "date_of_joining": me_["date_of_joining"],
        "referral_from": ref_id, "referral_from_name": _name(ref, ref_id),
        "position": position,
        "submitted_at": datetime.now(timezone.utc),
    })
    background.add_task(send_referral_email, me_["empid"], me_["name"], "existing", position,
                        _name(ref, ""))                             # HR acknowledgement mail (names only, no codes), after the response
    return {"success": True}


# ── 2. Referral with resume ───────────────────────────────────────────────────
@router.post("/referral")
async def submit_referral(
    background: BackgroundTasks,
    role: str = Form(...),
    resume: UploadFile = File(...),
    current_user: str = Depends(get_current_user),
):
    role = role.strip()
    if not role:
        raise HTTPException(400, "Role is required")
    filename = (resume.filename or "").strip()
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext not in _ALLOWED:
        raise HTTPException(400, "Only .pdf and .docx files are allowed")
    data = await resume.read(MAX_RESUME_BYTES + 1)
    if not data:
        raise HTTPException(400, "The uploaded file is empty")
    if len(data) > MAX_RESUME_BYTES:
        raise HTTPException(400, "Resume must be 10 MB or smaller")
    mime, magic = _ALLOWED[ext]
    if not data.startswith(magic):
        raise HTTPException(400, f"The file content is not a valid {ext[1:].upper()} file")
    safe_name = re.sub(r"[^\w.\- ]", "_", filename)[:120]
    file_id = _fs.put(data, filename=safe_name, content_type=mime)
    me_ = _profile(_emp(current_user), current_user.strip().upper())
    resume_referrals.insert_one({
        "submitted_by": me_["empid"], "submitted_by_name": me_["name"],
        "date_of_joining": me_["date_of_joining"],
        "role": role, "resume_file_id": file_id, "resume_filename": safe_name,
        "resume_size": len(data), "submitted_at": datetime.now(timezone.utc),
    })
    background.add_task(send_referral_email, me_["empid"], me_["name"], "referral", role)   # HR acknowledgement mail
    return {"success": True}


# ── Admin (control panel: Buddy Referral) ─────────────────────────────────────
@router.get("/admin/access-check")
def admin_access_check(current_user: str = Depends(get_current_user)):
    return {"is_admin": is_referral_admin(current_user)}


def _existing_rows():
    docs = list(existing_referrals.find({}).sort("submitted_at", -1))
    prof = _profiles_for(d.get("submitted_by") for d in docs)
    rows = []
    for d in docs:
        p = prof.get((d.get("submitted_by") or "").strip().upper()) or _profile(None, d.get("submitted_by") or "")
        rows.append({
            "id": str(d["_id"]),
            "empid": d.get("submitted_by"), "name": p["name"] or d.get("submitted_by_name"),
            "date_of_joining": p["date_of_joining"] or d.get("date_of_joining") or "",
            "manager": p["manager"], "partner": p["partner"],
            "referral_from": d.get("referral_from"), "referral_from_name": d.get("referral_from_name"),
            "position": d.get("position"), "submitted_at": d.get("submitted_at"),
        })
    return rows


def _referral_rows():
    docs = list(resume_referrals.find({}).sort("submitted_at", -1))
    prof = _profiles_for(d.get("submitted_by") for d in docs)
    rows = []
    for d in docs:
        p = prof.get((d.get("submitted_by") or "").strip().upper()) or _profile(None, d.get("submitted_by") or "")
        rows.append({
            "id": str(d["_id"]),
            "empid": d.get("submitted_by"), "name": p["name"] or d.get("submitted_by_name"),
            "date_of_joining": p["date_of_joining"] or d.get("date_of_joining") or "",
            "manager": p["manager"], "partner": p["partner"],
            "role": d.get("role"), "resume_filename": d.get("resume_filename"),
            "resume_size": d.get("resume_size"), "submitted_at": d.get("submitted_at"),
        })
    return rows


@router.get("/admin/existing")
def admin_existing(_: str = Depends(require_referral_admin)):
    rows = _existing_rows()
    for r in rows:
        r["submitted_at"] = _iso(r["submitted_at"])
    return {"rows": rows}


@router.get("/admin/referrals")
def admin_referrals(_: str = Depends(require_referral_admin)):
    rows = _referral_rows()
    for r in rows:
        r["submitted_at"] = _iso(r["submitted_at"])
    return {"rows": rows}


@router.get("/admin/referrals/{referral_id}/resume")
def admin_resume(referral_id: str, inline: int = 0, _: str = Depends(require_referral_admin)):
    """Resume bytes. inline=1 -> for viewing in the browser, else a download."""
    try:
        doc = resume_referrals.find_one({"_id": ObjectId(referral_id)})
    except (InvalidId, TypeError):
        raise HTTPException(404, "Referral not found")
    if not doc:
        raise HTTPException(404, "Referral not found")
    try:
        f = _fs.get(doc["resume_file_id"])
    except NoFile:
        raise HTTPException(404, "Resume file is missing")
    name = doc.get("resume_filename") or "resume"
    return Response(
        content=f.read(), media_type=getattr(f, "content_type", None) or "application/octet-stream",
        headers={"Content-Disposition": f'{"inline" if inline else "attachment"}; filename="{name}"'},
    )


def _xlsx(title: str, headers, rows) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2C3E50")
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([("" if v is None else v) for v in r])
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None] or [10])
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(width + 2, 12), 45)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@router.get("/admin/existing/export")
def export_existing(_: str = Depends(require_referral_admin)):
    rows = _existing_rows()
    data = _xlsx("Existing Employee Referrals",
                 ["Emp Code", "Emp Name", "Date of Joining", "Manager", "Partner",
                  "Referral From (Code)", "Referral From (Name)", "Position", "Submitted On"],
                 [[r["empid"], r["name"], r["date_of_joining"], r["manager"], r["partner"],
                   r["referral_from"], r["referral_from_name"], r["position"], _ist_text(r["submitted_at"])] for r in rows])
    return Response(content=data, media_type=_XLSX,
                    headers={"Content-Disposition": "attachment; filename=Existing_Employee_Referrals.xlsx"})


@router.get("/admin/referrals/export")
def export_referrals(_: str = Depends(require_referral_admin)):
    rows = _referral_rows()
    data = _xlsx("Referrals",
                 ["Emp Code", "Emp Name", "Date of Joining", "Manager", "Partner", "Role", "Resume File", "Submitted On"],
                 [[r["empid"], r["name"], r["date_of_joining"], r["manager"], r["partner"],
                   r["role"], r["resume_filename"], _ist_text(r["submitted_at"])] for r in rows])
    return Response(content=data, media_type=_XLSX,
                    headers={"Content-Disposition": "attachment; filename=Referrals.xlsx"})
